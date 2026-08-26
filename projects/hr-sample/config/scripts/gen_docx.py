# -*- coding: utf-8 -*-
"""docs 샘플 (DOCX) 3종 — 표준근로계약서 템플릿, 직무기술서(3직무), 인사팀 업무 매뉴얼(SOP).
vault에 docx 변환 레인이 없어 pptx와 동일하게 임시(unofficial) 레인으로 취급한다."""
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import data as D

OUT = os.path.expanduser("~/Documents/hr-samples/")
C = D.COMPANY

def base_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "맑은 고딕"
    style.font.size = Pt(10.5)
    return doc

def h1(doc, text):
    p = doc.add_heading(level=1)
    r = p.add_run(text); r.font.name = "맑은 고딕"; r.font.size = Pt(18)

def h2(doc, text):
    p = doc.add_heading(level=2)
    r = p.add_run(text); r.font.name = "맑은 고딕"; r.font.size = Pt(13)

def para(doc, text, italic=False, size=None, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.name = "맑은 고딕"
    if size: r.font.size = Pt(size)
    if italic: r.italic = True
    if color: r.font.color.rgb = RGBColor(*color)
    return p

def bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(text); r.font.name = "맑은 고딕"

def table(doc, header, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, htxt in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(htxt)
        r.bold = True; r.font.name = "맑은 고딕"; r.font.size = Pt(9.5)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(val))
            r.font.name = "맑은 고딕"; r.font.size = Pt(9.5)
    return t

def fixture_note(doc):
    p = para(doc, D.FIXTURE_NOTE, italic=True, size=9, color=(0x88,0x88,0x88))


# ══════════════════════════════════════════════════════════════════
# 1. 표준근로계약서 템플릿
# ══════════════════════════════════════════════════════════════════
def contract_template():
    doc = base_doc()
    h1(doc, "%s 표준근로계약서" % C)
    para(doc, "문서번호: HR-FORM-001 · 근거: 「취업규칙」(HR-REG-002) 제11조 · 개정 2025-07-01")
    fixture_note(doc)
    doc.add_paragraph()

    para(doc, "%s(이하 \"회사\"라 한다)과(와) ______________(이하 \"근로자\"라 한다)은 다음과 "
              "같이 근로계약을 체결한다." % C)

    h2(doc, "1. 근로계약기간")
    bullet(doc, "입사일: 20____년 ____월 ____일")
    bullet(doc, "계약기간: 기간의 정함이 없음 (단, 수습기간 3개월 별도 적용 — 「취업규칙」 제12조)")

    h2(doc, "2. 근무 장소 및 업무 내용")
    bullet(doc, "근무 장소: 본사 (재택근무는 「복무규정」 제3조에 따라 별도 신청)")
    bullet(doc, "담당 업무: ______________ (직무기술서 별첨)")
    bullet(doc, "소속 본부: ______________")

    h2(doc, "3. 직급 및 근로조건")
    table(doc, ["항목", "내용"], [
        ["직급", "L1 (사원) — 「급여규정」(HR-REG-003) 제3조"],
        ["근로시간", "09:00 ~ 18:00 (휴게 1시간) — 「취업규칙」 제8조·제9조"],
        ["근무일", "주 5일 (월~금)"],
        ["시차출퇴근", "08:00~10:00 사이 신청 가능 — 「복무규정」 제2조"],
    ])

    h2(doc, "4. 임금")
    table(doc, ["항목", "내용"], [
        ["기본급(연봉)", "금 ______________ 원 (「급여규정」 제3조 직급별 범위 내)"],
        ["지급일", "매월 25일 — 「급여규정」 제6조"],
        ["상여금", "평가 등급별 기본급 대비 0~200% — 「급여규정」 제9조"],
        ["수습기간 임금", "정규 임금의 100% — 「취업규칙」 제12조 제3항 (상여 제외)"],
    ])

    h2(doc, "5. 연차유급휴가")
    para(doc, "근속연수별 부여일수는 「복무규정」(HR-REG-004) 제5조 별표에 따른다. 입사 첫해는 "
              "월 1일씩 발생하며 최대 11일을 한도로 한다.")

    h2(doc, "6. 수습평가 및 정규 전환")
    para(doc, "수습기간(3개월) 종료 2주 전 수습평가를 실시하며, 기준은 입사 첫 주에 확정하는 "
              "30-60-90일 목표의 달성 여부다. 기준 미달 시 1회에 한해 1개월 연장할 수 있다 — "
              "「취업규칙」 제13조.")

    h2(doc, "7. 비밀유지 및 겸직금지")
    para(doc, "근로자는 재직 중 및 퇴직 후에도 업무상 알게 된 회사의 영업비밀을 누설하여서는 "
              "아니 되며, 회사의 사전 승인 없이 타 업체에 재직하거나 사업을 영위할 수 없다 — "
              "「취업규칙」 제16조·제17조.")

    doc.add_paragraph()
    para(doc, "본 계약의 내용에 대해 회사와 근로자는 각 1부씩 보관한다.")
    doc.add_paragraph()
    table(doc, ["구분", "서명"], [["회사 (대표이사)", "(인)"], ["근로자", "(인)"]])

    p = OUT + "나린테크 표준근로계약서(템플릿).docx"
    doc.save(p)
    return p, os.path.getsize(p)


# ══════════════════════════════════════════════════════════════════
# 2. 직무기술서 3종
# ══════════════════════════════════════════════════════════════════
JOBS = [
 {
  "code": "JD-ENG-002", "title": "백엔드 개발자 (선임)", "grade": "L2", "dept": "플랫폼개발본부",
  "summary": "사내 핵심 서비스의 API 서버와 데이터 파이프라인을 설계·구현·운영한다. "
             "주니어 개발자 1~2인의 코드리뷰와 기술 멘토링을 겸한다.",
  "responsibilities": [
    "REST/GraphQL API 설계 및 구현",
    "데이터베이스 스키마 설계와 마이그레이션 관리",
    "서비스 장애 대응 및 온콜 로테이션 참여",
    "주니어 개발자 코드리뷰 및 기술 멘토링",
    "분기 단위 기술 부채 정리 계획 수립",
  ],
  "requirements": [
    "백엔드 개발 경력 3년 이상",
    "Node.js 또는 Python 기반 서버 개발 경험",
    "관계형 데이터베이스 설계 및 쿼리 최적화 경험",
    "CI/CD 파이프라인 운영 경험",
  ],
  "preferred": ["대규모 트래픽 서비스 운영 경험", "클라우드 인프라(AWS/GCP) 운영 경험"],
  "okr_example": ["신규 결제 API 응답속도 P95 300ms 이하 달성", "레거시 배치 작업 2건을 이벤트 기반으로 전환"],
 },
 {
  "code": "JD-HR-002", "title": "인사담당 (선임)", "grade": "L2", "dept": "경영지원본부",
  "summary": "채용, 평가 운영, 온보딩·오프보딩 프로세스를 담당하며 인사 규정의 개정 실무를 "
             "지원한다.",
  "responsibilities": [
    "채용 공고 게시, 서류·면접 일정 조율, 합격자 온보딩 준비",
    "반기 인사평가 운영 — 일정 공지, 시스템 오픈, 보정회의 지원",
    "이의신청 접수 및 심의 일정 관리",
    "「인사평가 운영지침」·「급여규정」·「복무규정」 개정안 초안 작성 지원",
    "인사위원회 회의록 작성 및 결정 사항 실행",
  ],
  "requirements": [
    "인사 실무 경력 2년 이상",
    "인사평가·급여 프로세스 운영 경험",
    "엑셀/스프레드시트 기반 인력 데이터 관리 능력",
  ],
  "preferred": ["HRIS(인사정보시스템) 운영 경험", "노무 관련 기초 지식"],
  "okr_example": ["평가 결과 통보~이의신청 처리 평균 소요일 20% 단축", "온보딩 만족도 서베이 4.0/5.0 이상 유지"],
 },
 {
  "code": "JD-PM-003", "title": "프로덕트 매니저 (책임)", "grade": "L3", "dept": "사업본부",
  "summary": "제품 로드맵을 수립하고 개발·디자인·사업 조직 간 우선순위를 조율한다. 분기 "
             "OKR 수립과 달성도 추적을 주도한다.",
  "responsibilities": [
    "분기 OKR 수립 및 대상기간 개시 시점 확정 — 확정 기한 준수 책임",
    "제품 백로그 우선순위 관리 및 스프린트 계획 참여",
    "고객 피드백 수집·분석 및 제품 개선 반영",
    "타 본부(개발·디자인·사업) 간 커뮤니케이션 및 협업 조율",
  ],
  "requirements": [
    "프로덕트 매니저 경력 4년 이상",
    "B2B 또는 B2C SaaS 제품 운영 경험",
    "데이터 기반 의사결정 경험 (SQL 활용 가능)",
  ],
  "preferred": ["팀 리딩 경험", "OKR 운영 경험"],
  "okr_example": ["신규 기능 3건 출시 및 활성 사용자 지표 15% 개선", "분기 OKR 대상기간 개시 전 100% 확정"],
 },
]

def jd_doc(job):
    doc = base_doc()
    h1(doc, "직무기술서 — %s" % job["title"])
    para(doc, "문서번호: %s · 소속: %s · 직급: %s · 작성 2026-02-01" %
              (job["code"], job["dept"], job["grade"]))
    fixture_note(doc)
    doc.add_paragraph()

    h2(doc, "직무 요약")
    para(doc, job["summary"])

    h2(doc, "주요 업무")
    for r in job["responsibilities"]:
        bullet(doc, r)

    h2(doc, "자격 요건")
    for r in job["requirements"]:
        bullet(doc, r)

    h2(doc, "우대 사항")
    for r in job["preferred"]:
        bullet(doc, r)

    h2(doc, "직급 및 처우 (참고)")
    lo, hi, std = next((l, hi_, s) for code, name, l, hi_, s in D.GRADE_LEVELS if code == job["grade"])
    table(doc, ["항목", "내용"], [
        ["직급", "%s (%s)" % (job["grade"],
            next(name for code, name, *_ in D.GRADE_LEVELS if code == job["grade"]))],
        ["연봉 범위", "%d만원 ~ %d만원 (표준 %d만원)" % (lo, hi, std)],
        ["근거", "「급여규정」(HR-REG-003) 제3조"],
    ])

    h2(doc, "평가 시 OKR 예시")
    para(doc, "아래는 채용 공고 참고용 예시이며, 실제 OKR은 입사 후 팀 목표에 맞춰 협의한다.")
    for r in job["okr_example"]:
        bullet(doc, r)

    p = OUT + "직무기술서_%s.docx" % job["title"].split(" ")[0]
    doc.save(p)
    return p, os.path.getsize(p)


# ══════════════════════════════════════════════════════════════════
# 3. 인사팀 업무 매뉴얼 (SOP)
# ══════════════════════════════════════════════════════════════════
def sop_doc():
    doc = base_doc()
    h1(doc, "%s 인사팀 업무 매뉴얼 (SOP)" % C)
    para(doc, "문서번호: HR-SOP-001 · 작성 2026-02-15 · 대상: 인사담당 신규 인력 온보딩용")
    fixture_note(doc)
    doc.add_paragraph()

    para(doc, "이 매뉴얼은 인사담당 부서의 반복 업무를 절차화한 것이다. 개별 규정의 상세는 "
              "각 규정 문서를 참조하고, 이 문서는 \"언제 무엇을 하는가\"의 실행 순서에 집중한다.")

    h2(doc, "1. 반기 평가 운영 타임라인")
    table(doc, ["시점", "작업", "근거 조문"], [
        ["대상기간 개시일", "OKR 확정 마감 공지", "지침 제5조 제3항"],
        ["대상기간 종료 + 1주", "자기평가 오픈 공지", "지침 제9조 제2항(2024) / 관행상 유지"],
        ["대상기간 종료 + 3주", "1차 평가 마감 확인", "지침 제9조 제3항"],
        ["대상기간 종료 + 4주", "2차 평가 마감 확인", "지침 제9조 제4항"],
        ["2차 평가 마감 + 1주", "보정회의 일정 조율·개최", "지침 제11조"],
        ["보정회의 종료 + 3일", "결과 확정 및 통보 발송 (등급 근거 포함)", "지침 제12조 제1항"],
        ["통보일 + 7일", "이의신청 접수 마감", "지침 제12조 제2항"],
        ["접수 마감 + 14일", "이의신청 심의 완료·통보", "지침 제12조 제3항"],
        ["결과 확정 익월", "연봉 인상·상여 반영 급여 지급", "급여규정 제7조·제9조"],
    ])

    h2(doc, "2. 채용 프로세스 체크리스트")
    for step in [
        "채용 요청서 접수 — 본부장 승인 확인",
        "채용 공고 게시 (사내 채용공고 페이지 + 외부 채널)",
        "서류전형 — 접수 마감 후 3영업일 이내 1차 스크리닝",
        "직무면접 일정 조율 — 현업 면접관 2인 이상 배정",
        "임원면접 — 최종 후보자 대상",
        "합격 통보 및 처우 협의 — 「급여규정」 제3조 직급별 범위 내에서 결정",
        "근로계약서 작성 — 「표준근로계약서」(HR-FORM-001) 사용",
        "입사 D-3 — 장비·계정 신청, 멘토 배정 요청",
    ]:
        bullet(doc, step)

    h2(doc, "3. 온보딩 지원 체크리스트")
    for step in [
        "입사 당일 — 오리엔테이션(급여·복리후생·근태 안내)",
        "입사 D+3 — 멘토·신규입사자의 30-60-90일 목표 초안 확인",
        "입사 30일 — 목표 진행 상황 체크인",
        "입사 60일 — 목표 진행 상황 체크인",
        "수습종료 D-14 — 수습평가 실시 안내",
        "수습종료 D-7 — 수습평가 결과 취합, 필요 시 연장 절차 안내",
        "수습종료일 — 정규 전환 통보 또는 연장·미전환 사유 서면 통지",
    ]:
        bullet(doc, step)

    h2(doc, "4. 퇴사(오프보딩) 체크리스트")
    for step in [
        "퇴사 통보 접수 — 사유 구분(자발/비자발) 기록",
        "인수인계 계획서 제출 요청 (최소 2주 전)",
        "최종 근무일 확인 및 급여·연차 정산 계산",
        "계정·장비 회수 일정 조율",
        "퇴직금 산정 — 근로자퇴직급여보장법 기준 (급여규정 제12조)",
        "퇴사 사유 분류 결과를 분기 채용·이직 리포트에 반영",
    ]:
        bullet(doc, step)

    h2(doc, "5. 문서 체계 한눈에 보기")
    table(doc, ["문서번호", "문서명", "성격"], [
        ["HR-REG-001", "인사평가 운영지침", "평가 절차·등급·보정회의·이의신청"],
        ["HR-REG-002", "취업규칙 (발췌)", "근로시간·채용수습·복무·휴가"],
        ["HR-REG-003", "급여규정", "직급·인상률·상여·성과급"],
        ["HR-REG-004", "복무규정", "근태·연차·경조사·병가"],
        ["HR-FORM-001", "표준근로계약서", "채용 확정자용 계약 템플릿"],
        ["HR-SOP-001", "인사팀 업무 매뉴얼", "이 문서 — 실행 절차"],
    ])

    h2(doc, "6. 자주 참조하는 산출물")
    for step in [
        "반기 인사평가 결과 리포트 — 평가 차수 종료 후 인사위원회 보고용",
        "연간 인사운영 보고서 — 연 1회, 익년 초 작성",
        "분기 채용·이직 리포트 — 매 분기 종료 후 작성",
        "인사위원회 회의록 — 안건별 결정 사항의 유일한 근거",
    ]:
        bullet(doc, step)

    p = OUT + "나린테크 인사팀 업무 매뉴얼(SOP).docx"
    doc.save(p)
    return p, os.path.getsize(p)


if __name__ == "__main__":
    made = []
    p, sz = contract_template(); made.append((os.path.basename(p), sz))
    for job in JOBS:
        p, sz = jd_doc(job); made.append((os.path.basename(p), sz))
    p, sz = sop_doc(); made.append((os.path.basename(p), sz))
    for n, sz in made:
        print("  %6d B  %s" % (sz, n))
    print("DOCX %d종 생성" % len(made))
