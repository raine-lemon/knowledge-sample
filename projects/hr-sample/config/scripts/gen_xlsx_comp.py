# -*- coding: utf-8 -*-
"""「나린테크 급여근태 대장(2026H1)」 XLSX 생성기.

data.py(단일 진실원) + roster.py(개인 단위 결정론적 로스터)에서 급여·근태 데이터를
파생해 시트 7개짜리 워크북을 만든다. 임의로 지어낸 수치는 없다 — 모든 값은
소스 모듈의 함수·상수에서 계산되거나, 사번 기반 고정 시드 난수로 결정론적으로
파생된다(§ 연차 사용일수 파생).

산출물: ~/Documents/hr-samples/나린테크 급여근태 대장(2026H1).xlsx

실행:
    cd <repo>/projects/hr-sample/config/scripts
    <python> gen_xlsx_comp.py

주의: data.py / roster.py 는 읽기 전용으로만 쓴다. 이 스크립트는 두 모듈을
      절대 수정하지 않는다.
"""
import hashlib
import os
import random
import statistics
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import data as D
import roster as R

# ── 기준일 ────────────────────────────────────────────────────────
# 2026년 상반기(2026-H1) 말일. roster.py의 파생 함수 기본값과 동일하게 맞춘다.
ASOF = date(2026, 6, 30)
HALF_KEY = (2026, 1)          # 2026 상반기 평가 차수 키 (p["grades"]의 키 형식)

OUT_DIR = os.path.expanduser("~/Documents/hr-samples")
OUT_PATH = os.path.join(OUT_DIR, "나린테크 급여근태 대장(2026H1).xlsx")

# ── 서식 상수 ─────────────────────────────────────────────────────
FMT_MONEY = "#,##0"           # 금액(만원)
FMT_PCT = '0.0"%"'            # 백분율 — 값 자체가 8.0 같은 퍼센트 숫자다
FMT_DAY = "0.0"               # 일수(소수 허용)
FMT_YEAR = "0.0"              # 근속연수

HEAD_FILL = PatternFill("solid", fgColor="1F3864")     # 헤더 배경(진한 남색)
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(size=10)
TOTAL_FONT = Font(bold=True, size=10)
TOP_BORDER = Border(top=Side(style="thin", color="1F3864"))
CENTER = Alignment(horizontal="center", vertical="center")


# ══════════════════════════════════════════════════════════════════
# § 연차 사용일수 파생 (roster.py에 없는 값 — 결정론적으로 만든다)
# ══════════════════════════════════════════════════════════════════
def _stable_seed(emp_id):
    """사번 문자열 → 고정 정수 시드.

    파이썬 내장 hash()는 PYTHONHASHSEED에 따라 프로세스마다 달라져 재현성이
    깨진다. 그래서 SHA-256 해시 앞 8바이트를 정수로 쓴다 — 어떤 머신·어떤
    실행에서도 같은 사번이면 같은 시드다.
    """
    h = hashlib.sha256(emp_id.encode("utf-8")).hexdigest()
    return int(h[:16], 16)


def leave_used(p, entitled):
    """연차 사용일수(일). 산식은 README 시트에 그대로 기록된다.

        rng   = random.Random(int(sha256(사번)[:16], 16))
        비율   = clamp(rng.gauss(0.68, 0.18), 0.05, 1.00)
        사용일 = round(부여일수 × 비율)  →  [0, 부여일수]로 클램프

    평균 소진율 0.68은 「복무규정」 제7조 제2항의 "반기 잔여 연차 50% 이상
    소진 권장"을 상회하는 현실적인 수준으로 잡은 목업 파라미터다.
    반차·반반차(제6조)는 반영하지 않고 1일 단위 정수로만 만든다.
    """
    rng = random.Random(_stable_seed(p["emp_id"]))
    ratio = rng.gauss(0.68, 0.18)
    ratio = max(0.05, min(1.00, ratio))
    used = int(round(entitled * ratio))
    return max(0, min(entitled, used))


# ══════════════════════════════════════════════════════════════════
# § 레코드 구성
# ══════════════════════════════════════════════════════════════════
LEVEL_NAME = {c: n for c, n, _lo, _hi, _std in D.GRADE_LEVELS}
LEVEL_BAND = {c: (lo, hi) for c, _n, lo, hi, _std in D.GRADE_LEVELS}
LEVEL_ORDER = [c for c, _n, _lo, _hi, _std in D.GRADE_LEVELS]


def build_records(active):
    """재직자 1인 = 1레코드. 급여·근태 파생값을 모두 계산해 둔다."""
    recs = []
    for p in sorted(active, key=lambda x: x["emp_id"]):
        lvl = R.level_of(p, ASOF)
        tenure = R.tenure_years(p, ASOF)
        annual = R.salary(p, ASOF)                    # 연봉(만원) — 밴드 클램프 포함
        monthly = annual / 12.0                        # 월기본급(만원)
        grade = p["grades"].get(HALF_KEY)               # 2026H1 등급 (없으면 평가제외자)

        if grade:
            raise_pct = D.GRADE_TO_RAISE[grade]
            bonus_pct = D.GRADE_TO_BONUS[grade]
            bonus = monthly * bonus_pct / 100.0         # 상여금 = 월기본급 × 지급률
        else:
            # 2026-Q2 입사자 = 수습 중. 「인사평가 운영지침」 제10조로 정규 평가에서
            # 분리되고, 「급여규정」 제9조 ④에 따라 상여금도 지급하지 않는다.
            raise_pct = bonus_pct = None
            bonus = 0.0

        entitled = R.leave_entitled(p, ASOF)
        used = leave_used(p, entitled)

        recs.append({
            "emp_id": p["emp_id"],
            "dept": p["dept"],
            "level": lvl,
            "level_label": "%s(%s)" % (lvl, LEVEL_NAME[lvl]),
            "tenure": tenure,
            "grade": grade,
            "annual": annual,
            "monthly": monthly,
            "raise_pct": raise_pct,
            "bonus_pct": bonus_pct,
            "bonus": bonus,
            "total_comp": annual + bonus,
            "entitled": entitled,
            "used": used,
            "remain": entitled - used,
            "use_rate": used / entitled * 100.0 if entitled else 0.0,
        })
    return recs


# ══════════════════════════════════════════════════════════════════
# § 시트 작성 헬퍼
# ══════════════════════════════════════════════════════════════════
def write_header(ws, headers, row=1):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = CENTER
    ws.row_dimensions[row].height = 22


def set_widths(ws, widths, start=1):
    for i, w in enumerate(widths, start=start):
        ws.column_dimensions[get_column_letter(i)].width = w


def style_total_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = TOTAL_FONT
        cell.border = TOP_BORDER


def apply_formats(ws, first_row, last_row, col_fmt):
    """{열번호: 표시서식} 을 데이터 구간에 일괄 적용."""
    for r in range(first_row, last_row + 1):
        for c, fmt in col_fmt.items():
            ws.cell(row=r, column=c).number_format = fmt
            if ws.cell(row=r, column=c).font is None or not ws.cell(row=r, column=c).font.bold:
                ws.cell(row=r, column=c).font = BODY_FONT


# ══════════════════════════════════════════════════════════════════
# 시트 1: 급여대장
# ══════════════════════════════════════════════════════════════════
def sheet_payroll(wb, recs):
    ws = wb.create_sheet("급여대장")
    headers = ["사번", "본부", "직급", "근속연수", "2026H1등급", "연봉(만원)",
               "월기본급(만원)", "인상률(%)", "상여지급률(%)", "상여금(만원)",
               "연간총보상(만원)"]
    write_header(ws, headers)

    for i, r in enumerate(recs, start=2):
        ws.cell(row=i, column=1, value=r["emp_id"])
        ws.cell(row=i, column=2, value=r["dept"])
        ws.cell(row=i, column=3, value=r["level_label"])
        ws.cell(row=i, column=4, value=r["tenure"])
        ws.cell(row=i, column=5, value=r["grade"] if r["grade"] else None)
        ws.cell(row=i, column=6, value=r["annual"])
        ws.cell(row=i, column=7, value=round(r["monthly"], 1))
        ws.cell(row=i, column=8, value=r["raise_pct"])          # 평가제외자는 None → 빈칸
        ws.cell(row=i, column=9, value=r["bonus_pct"])          # 평가제외자는 None → 빈칸
        ws.cell(row=i, column=10, value=round(r["bonus"], 1) if r["grade"] else None)
        ws.cell(row=i, column=11, value=round(r["total_comp"], 1))

    last = len(recs) + 1
    apply_formats(ws, 2, last, {
        4: FMT_YEAR, 6: FMT_MONEY, 7: FMT_MONEY, 8: FMT_PCT, 9: FMT_PCT,
        10: FMT_MONEY, 11: FMT_MONEY,
    })
    for r in range(2, last + 1):
        ws.cell(row=r, column=5).alignment = CENTER
    set_widths(ws, [10, 16, 12, 10, 12, 12, 14, 11, 14, 13, 15])
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), last)
    ws.freeze_panes = "A2"
    return ws


# ══════════════════════════════════════════════════════════════════
# 시트 2: 직급별급여분포
# ══════════════════════════════════════════════════════════════════
def sheet_band(wb, recs):
    ws = wb.create_sheet("직급별급여분포")
    headers = ["직급", "인원", "규정하한(만원)", "규정상한(만원)", "실제최소(만원)",
               "실제최대(만원)", "중위값(만원)", "평균(만원)", "밴드이탈인원"]
    write_header(ws, headers)

    rows = []
    row = 2
    for lvl in LEVEL_ORDER:
        grp = [r for r in recs if r["level"] == lvl]
        lo, hi = LEVEL_BAND[lvl]
        if grp:
            sal = [r["annual"] for r in grp]
            out = sum(1 for s in sal if s < lo or s > hi)   # 밴드이탈 — salary()가 클램프하므로 0
            vals = [len(grp), lo, hi, min(sal), max(sal),
                    round(statistics.median(sal), 1), round(statistics.mean(sal), 1), out]
        else:
            vals = [0, lo, hi, None, None, None, None, 0]
        ws.cell(row=row, column=1, value="%s(%s)" % (lvl, LEVEL_NAME[lvl]))
        for c, v in enumerate(vals, start=2):
            ws.cell(row=row, column=c, value=v)
        rows.append((lvl, vals))
        row += 1

    data_last = row - 1
    # 합계 행
    ws.cell(row=row, column=1, value="합계")
    ws.cell(row=row, column=2, value=sum(v[0] for _l, v in rows))
    all_sal = [r["annual"] for r in recs]
    ws.cell(row=row, column=5, value=min(all_sal))
    ws.cell(row=row, column=6, value=max(all_sal))
    ws.cell(row=row, column=7, value=round(statistics.median(all_sal), 1))
    ws.cell(row=row, column=8, value=round(statistics.mean(all_sal), 1))
    ws.cell(row=row, column=9, value=sum(v[7] for _l, v in rows))
    style_total_row(ws, row, len(headers))
    total_row = row

    apply_formats(ws, 2, data_last, {c: FMT_MONEY for c in (3, 4, 5, 6, 7, 8)})
    for c in (3, 4, 5, 6, 7, 8):
        ws.cell(row=total_row, column=c).number_format = FMT_MONEY
    set_widths(ws, [14, 8, 15, 15, 15, 15, 14, 13, 14])
    ws.freeze_panes = "A2"

    # ── BarChart: 직급별 중위값·평균 연봉 ──
    ch = BarChart()
    ch.type = "col"
    ch.style = 10
    ch.title = "직급별 연봉 중위값·평균 (만원)"
    ch.y_axis.title = "연봉(만원)"
    ch.x_axis.title = "직급"
    dat = Reference(ws, min_col=7, max_col=8, min_row=1, max_row=data_last)
    cats = Reference(ws, min_col=1, min_row=2, max_row=data_last)
    ch.add_data(dat, titles_from_data=True)
    ch.set_categories(cats)
    ch.height, ch.width = 8, 18
    ws.add_chart(ch, "K2")
    return ws


# ══════════════════════════════════════════════════════════════════
# 시트 3: 등급별처우
# ══════════════════════════════════════════════════════════════════
def sheet_grade(wb, recs, hf):
    ws = wb.create_sheet("등급별처우")
    headers = ["등급", "인원", "비율(%)", "인상률(%)", "상여지급률(%)",
               "평균연봉(만원)", "평균상여(만원)", "인건비영향(만원)"]
    write_header(ws, headers)

    graded = [r for r in recs if r["grade"]]
    n = len(graded)
    row = 2
    for g in D.GRADES:
        grp = [r for r in graded if r["grade"] == g]
        ws.cell(row=row, column=1, value=g)
        ws.cell(row=row, column=2, value=len(grp))
        ws.cell(row=row, column=3, value=round(len(grp) / n * 100, 1))
        ws.cell(row=row, column=4, value=D.GRADE_TO_RAISE[g])
        ws.cell(row=row, column=5, value=D.GRADE_TO_BONUS[g])
        ws.cell(row=row, column=6,
                value=round(statistics.mean([r["annual"] for r in grp]), 1) if grp else None)
        ws.cell(row=row, column=7,
                value=round(statistics.mean([r["bonus"] for r in grp]), 1) if grp else None)
        # 인건비영향 = 해당 등급 인원의 상여금 합계
        ws.cell(row=row, column=8, value=round(sum(r["bonus"] for r in grp), 1))
        row += 1
    data_last = row - 1

    # 합계 행 — 가중평균 인상률·상여율은 data.py의 weighted_* 와 대조 검증한다
    w_raise = sum(len([r for r in graded if r["grade"] == g]) * D.GRADE_TO_RAISE[g]
                  for g in D.GRADES) / n
    w_bonus = sum(len([r for r in graded if r["grade"] == g]) * D.GRADE_TO_BONUS[g]
                  for g in D.GRADES) / n
    ws.cell(row=row, column=1, value="합계/가중평균")
    ws.cell(row=row, column=2, value=n)
    ws.cell(row=row, column=3, value=100.0)
    ws.cell(row=row, column=4, value=round(w_raise, 2))
    ws.cell(row=row, column=5, value=round(w_bonus, 2))
    ws.cell(row=row, column=6, value=round(statistics.mean([r["annual"] for r in graded]), 1))
    ws.cell(row=row, column=7, value=round(statistics.mean([r["bonus"] for r in graded]), 1))
    ws.cell(row=row, column=8, value=round(sum(r["bonus"] for r in graded), 1))
    style_total_row(ws, row, len(headers))
    total_row = row

    apply_formats(ws, 2, data_last, {3: FMT_PCT, 4: FMT_PCT, 5: FMT_PCT,
                                     6: FMT_MONEY, 7: FMT_MONEY, 8: FMT_MONEY})
    for c, f in {3: FMT_PCT, 4: '0.00"%"', 5: '0.00"%"',
                 6: FMT_MONEY, 7: FMT_MONEY, 8: FMT_MONEY}.items():
        ws.cell(row=total_row, column=c).number_format = f
    ws.cell(row=total_row + 2, column=1,
            value="대조: data.py weighted_raise(half(2026,1))=%.2f%% / weighted_bonus=%.2f%%"
                  % (D.weighted_raise(hf), D.weighted_bonus(hf)))
    ws.cell(row=total_row + 2, column=1).font = Font(size=9, italic=True)
    set_widths(ws, [16, 8, 10, 11, 14, 14, 14, 16])
    ws.freeze_panes = "A2"

    # ── BarChart: 등급별 인원 ──
    ch = BarChart()
    ch.type = "col"
    ch.style = 12
    ch.title = "2026H1 등급별 인원"
    ch.y_axis.title = "인원(명)"
    ch.x_axis.title = "등급"
    dat = Reference(ws, min_col=2, max_col=2, min_row=1, max_row=data_last)
    cats = Reference(ws, min_col=1, min_row=2, max_row=data_last)
    ch.add_data(dat, titles_from_data=True)
    ch.set_categories(cats)
    ch.height, ch.width = 8, 14
    ws.add_chart(ch, "J2")

    ch2 = BarChart()
    ch2.type = "col"
    ch2.style = 10
    ch2.title = "등급별 인건비영향 — 상여금 합계 (만원)"
    ch2.y_axis.title = "만원"
    ch2.x_axis.title = "등급"
    dat2 = Reference(ws, min_col=8, max_col=8, min_row=1, max_row=data_last)
    ch2.add_data(dat2, titles_from_data=True)
    ch2.set_categories(cats)
    ch2.height, ch2.width = 8, 14
    ws.add_chart(ch2, "J20")
    return ws, w_raise, w_bonus


# ══════════════════════════════════════════════════════════════════
# 시트 4: 연차현황
# ══════════════════════════════════════════════════════════════════
def sheet_leave(wb, recs):
    ws = wb.create_sheet("연차현황")
    headers = ["사번", "본부", "직급", "근속연수", "부여일수", "사용일수",
               "잔여일수", "소진율(%)"]
    write_header(ws, headers)
    for i, r in enumerate(recs, start=2):
        ws.cell(row=i, column=1, value=r["emp_id"])
        ws.cell(row=i, column=2, value=r["dept"])
        ws.cell(row=i, column=3, value=r["level_label"])
        ws.cell(row=i, column=4, value=r["tenure"])
        ws.cell(row=i, column=5, value=r["entitled"])
        ws.cell(row=i, column=6, value=r["used"])
        ws.cell(row=i, column=7, value=r["remain"])
        ws.cell(row=i, column=8, value=round(r["use_rate"], 1))
    last = len(recs) + 1
    apply_formats(ws, 2, last, {4: FMT_YEAR, 5: "0", 6: "0", 7: "0", 8: FMT_PCT})
    set_widths(ws, [10, 16, 12, 10, 11, 11, 11, 12])
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(headers)), last)
    ws.freeze_panes = "A2"
    return ws


# ══════════════════════════════════════════════════════════════════
# 시트 5: 본부별근태요약
# ══════════════════════════════════════════════════════════════════
def sheet_dept_leave(wb, recs):
    ws = wb.create_sheet("본부별근태요약")
    headers = ["본부", "인원", "평균부여", "평균사용", "평균잔여", "평균소진율(%)",
               "소진율50%미만인원"]
    write_header(ws, headers)
    row = 2
    for dept, _share in D.DEPT_SHARE:
        grp = [r for r in recs if r["dept"] == dept]
        ws.cell(row=row, column=1, value=dept)
        ws.cell(row=row, column=2, value=len(grp))
        ws.cell(row=row, column=3, value=round(statistics.mean([r["entitled"] for r in grp]), 1))
        ws.cell(row=row, column=4, value=round(statistics.mean([r["used"] for r in grp]), 1))
        ws.cell(row=row, column=5, value=round(statistics.mean([r["remain"] for r in grp]), 1))
        ws.cell(row=row, column=6, value=round(statistics.mean([r["use_rate"] for r in grp]), 1))
        # 「복무규정」 제7조 제2항의 50% 소진 권장에 미달하는 인원
        ws.cell(row=row, column=7, value=sum(1 for r in grp if r["use_rate"] < 50.0))
        row += 1
    data_last = row - 1

    ws.cell(row=row, column=1, value="전사 합계/평균")
    ws.cell(row=row, column=2, value=len(recs))
    ws.cell(row=row, column=3, value=round(statistics.mean([r["entitled"] for r in recs]), 1))
    ws.cell(row=row, column=4, value=round(statistics.mean([r["used"] for r in recs]), 1))
    ws.cell(row=row, column=5, value=round(statistics.mean([r["remain"] for r in recs]), 1))
    ws.cell(row=row, column=6, value=round(statistics.mean([r["use_rate"] for r in recs]), 1))
    ws.cell(row=row, column=7, value=sum(1 for r in recs if r["use_rate"] < 50.0))
    style_total_row(ws, row, len(headers))

    apply_formats(ws, 2, data_last, {3: FMT_DAY, 4: FMT_DAY, 5: FMT_DAY, 6: FMT_PCT})
    for c, f in {3: FMT_DAY, 4: FMT_DAY, 5: FMT_DAY, 6: FMT_PCT}.items():
        ws.cell(row=row, column=c).number_format = f
    set_widths(ws, [18, 8, 11, 11, 11, 14, 18])
    ws.freeze_panes = "A2"
    return ws


# ══════════════════════════════════════════════════════════════════
# 시트 6: 인건비추이
# ══════════════════════════════════════════════════════════════════
def sheet_trend(wb, recs):
    """반기 차수별 인건비(상여) 추정.

    ※ 근사다. roster.salary()는 2026-06-30 기준으로만 정의되며 과거 시점의
      연봉 이력을 모델링하지 않는다. 따라서 모든 차수에 2026-06-30 재직자
      212명의 평균 월기본급을 상수로 적용한다. 산식은 README 시트 참조.
    """
    ws = wb.create_sheet("인건비추이")
    headers = ["평가차수", "평가대상(명)", "가중평균인상률(%)", "가중평균상여율(%)",
               "추정상여총액(만원)"]
    write_header(ws, headers)

    avg_monthly = statistics.mean([r["monthly"] for r in recs])
    row = 2
    for y, h, *_rest in D.HALVES:
        hf = D.half(y, h)
        wr = D.weighted_raise(hf)
        wb_ = D.weighted_bonus(hf)
        est = hf["n"] * avg_monthly * wb_ / 100.0
        ws.cell(row=row, column=1, value="%d-H%d" % (y, h))
        ws.cell(row=row, column=2, value=hf["n"])
        ws.cell(row=row, column=3, value=round(wr, 2))
        ws.cell(row=row, column=4, value=round(wb_, 2))
        ws.cell(row=row, column=5, value=round(est, 1))
        row += 1
    data_last = row - 1

    apply_formats(ws, 2, data_last, {3: '0.00"%"', 4: '0.00"%"', 5: FMT_MONEY})
    ws.cell(row=data_last + 2, column=1,
            value="근사 산식: 추정상여총액 = 평가대상 인원 × 평균 월기본급(%.1f만원, "
                  "2026-06-30 재직 212명 기준 상수) × 가중평균상여율 ÷ 100" % avg_monthly)
    ws.cell(row=data_last + 2, column=1).font = Font(size=9, italic=True)
    set_widths(ws, [14, 13, 18, 18, 18])
    ws.freeze_panes = "A2"
    return ws, avg_monthly


# ══════════════════════════════════════════════════════════════════
# 시트 7: README
# ══════════════════════════════════════════════════════════════════
def sheet_readme(wb, recs, hf, avg_monthly, n_excluded):
    ws = wb.create_sheet("README")
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 118

    def h1(text):
        # 항상 한 줄 띄우고 새 섹션을 연다 (제목 행을 덮어쓰지 않도록 max_row+2 고정)
        r = ws.max_row + 2
        c = ws.cell(row=r, column=1, value=text)
        c.font = Font(bold=True, size=12, color="1F3864")
        return r

    def kv(k, v):
        r = ws.max_row + 1
        a = ws.cell(row=r, column=1, value=k)
        a.font = Font(bold=True, size=10)
        a.alignment = Alignment(vertical="top")
        b = ws.cell(row=r, column=2, value=v)
        b.font = BODY_FONT
        b.alignment = Alignment(vertical="top", wrap_text=True)
        return r

    t = ws.cell(row=1, column=1, value="나린테크 급여근태 대장 (2026H1) — README")
    t.font = Font(bold=True, size=14, color="1F3864")

    h1("0. 고지")
    kv("가상 목업",
       "본 워크북은 지식관리 파이프라인 검증용 가상 목업 데이터다. 실존 기업·인물·급여와 "
       "일절 무관하며, 어떤 실제 인사 의사결정에도 사용할 수 없다. 회사명 「(주)나린테크」는 가공의 법인이다.")
    kv("기준일", "2026-06-30 (2026년 상반기 말). 재직자 %d명, 2026H1 평가대상 %d명."
       % (len(recs), hf["n"]))

    h1("1. 데이터 출처")
    kv("단일 진실원",
       "config/scripts/data.py — 분기별 인력 이동(QUARTERS), 반기 등급 분포(HALVES), "
       "직급별 연봉밴드(GRADE_LEVELS), 등급별 인상률(GRADE_TO_RAISE)·상여지급률(GRADE_TO_BONUS), "
       "근속연수별 연차(LEAVE_BY_TENURE/annual_leave_days), 본부 배분(DEPT_SHARE).")
    kv("개인 로스터",
       "config/scripts/roster.py — build(seed=20260826)이 (people, active)를 반환한다. "
       "같은 seed면 항상 같은 로스터가 나온다. 파생 함수: tenure_years(), level_of(), "
       "salary(), leave_entitled().")
    kv("연봉 산정",
       "roster.salary() = 직급 표준연봉 × Π(1 + 등급별 인상률/100 × 0.5) 를 직급 연봉밴드 "
       "[하한, 상한]으로 클램프한 값(만원). 밴드 클램프 때문에 시트2의 밴드이탈인원은 항상 0이다.")
    kv("규정 근거",
       "「급여규정(2026 개정)」 제3조(직급 체계·연봉밴드), 제6조(급여 지급일 매월 25일), "
       "제8조(등급별 인상률 S8.0/A5.0/B3.0/C1.0/D0.0%), 제9조(상여 지급률 S200/A150/B100/C50/D0%, "
       "④ 수습기간 중 상여 미지급). 「복무규정(2025 개정)」 제5조(근속연수별 연차일수), "
       "제6조(반차·반반차), 제7조 ②(반기 50% 소진 권장).")

    h1("2. 시트 구성")
    kv("① 급여대장",
       "재직자 %d명 1인 1행. 월기본급 = 연봉÷12, 상여금 = 월기본급 × 상여지급률, "
       "연간총보상 = 연봉 + 상여금. 자동필터 + 첫 행 틀고정." % len(recs))
    kv("② 직급별급여분포",
       "L1~L5 직급별 인원·규정밴드(data.py GRADE_LEVELS)·실제 최소/최대/중위/평균 연봉과 "
       "밴드이탈인원. 중위값·평균 막대그래프 포함.\n"
       "※ L5(본부장) 인원이 0인 것은 오류가 아니다. 직급은 roster.LEVEL_BY_TENURE에 따라 근속연수로만 "
       "결정되고 L5는 근속 14년 이상인데, 이 로스터의 재직자 최장 근속은 12.7년이다. "
       "규정 밴드 행은 비교 기준으로 남겨 두었다.")
    kv("③ 등급별처우",
       "2026H1 등급(S~D)별 인원·비율·인상률·상여지급률·평균연봉·평균상여와 "
       "인건비영향(= 해당 등급 인원의 상여금 합계). 마지막 행은 합계 + 가중평균 인상률/상여율이며 "
       "data.py의 weighted_raise()/weighted_bonus()와 소수 둘째 자리까지 대조 검증한다. 막대그래프 2종.")
    kv("④ 연차현황",
       "재직자 %d명 1인 1행. 부여일수는 roster.leave_entitled()(복무규정 제5조), "
       "사용일수는 아래 §3의 결정론적 파생. 자동필터 + 첫 행 틀고정." % len(recs))
    kv("⑤ 본부별근태요약",
       "본부 5개의 평균 부여·사용·잔여·소진율과 소진율 50% 미만 인원(복무규정 제7조 ② 권고 미달). "
       "마지막에 전사 합계/평균 행.")
    kv("⑥ 인건비추이", "2024H1~2026H1 5개 차수의 평가대상 인원·가중평균 인상률/상여율·추정상여총액.")
    kv("⑦ README", "이 시트.")

    h1("3. 추정·파생으로 만든 값 (실제 소스에 없는 값)")
    kv("연차 사용일수 (시트④·⑤)",
       "roster.py에 없는 값이므로 사번 기반 고정 시드로 결정론적으로 생성했다.\n"
       "  seed  = int(sha256(사번).hexdigest()[:16], 16)   ← 파이썬 내장 hash()는 실행마다 달라져 쓰지 않음\n"
       "  rng   = random.Random(seed)\n"
       "  비율   = clamp(rng.gauss(mu=0.68, sigma=0.18), 0.05, 1.00)\n"
       "  사용일수 = round(부여일수 × 비율) 을 [0, 부여일수] 로 클램프\n"
       "따라서 잔여일수 = 부여일수 − 사용일수 ≥ 0 이 항상 성립한다. 평균 0.68은 복무규정 제7조 ②의 "
       "50% 소진 권장을 상회하도록 잡은 목업 파라미터이며, 실측치가 아니다. "
       "반차·반반차(제6조)는 반영하지 않고 1일 단위 정수로만 만든다. 같은 사번이면 어떤 머신에서 "
       "몇 번을 다시 돌려도 같은 값이 나온다.")
    kv("추정상여총액 (시트⑥)",
       "근사값이다. roster.salary()는 2026-06-30 기준으로만 정의되고 과거 차수의 연봉 이력을 "
       "모델링하지 않기 때문이다.\n"
       "  추정상여총액 = 해당 차수 평가대상 인원(data.half(y,h)[\"n\"]) × 평균 월기본급 × 가중평균상여율 ÷ 100\n"
       "  평균 월기본급 = %.2f만원 — 2026-06-30 재직 %d명의 (연봉÷12) 평균을 모든 차수에 상수로 적용\n"
       "실제로는 과거 차수의 평균 월기본급이 더 낮았을 것이므로 과거 차수 금액은 과대 추정이다. "
       "차수 간 비교 지표가 아니라 규모 감각용으로만 쓸 것." % (avg_monthly, len(recs)))
    kv("평가제외자 처리 (시트①)",
       "2026-Q2 입사자 %d명은 2026H1 평가 대상이 아니다 — 「인사평가 운영지침」 제10조에 따라 "
       "수습사원은 정규 평가에서 분리되고, 「급여규정」 제9조 ④에 따라 상여금도 지급하지 않는다. "
       "이들의 2026H1등급·인상률·상여지급률·상여금은 빈칸이며, 연간총보상 = 연봉이다. "
       "시트③(등급별처우)의 인원 합계 %d명은 이 %d명을 제외한 값이다."
       % (n_excluded, hf["n"], n_excluded))

    h1("4. 재현 방법")
    kv("실행",
       "cd <repo>/projects/hr-sample/config/scripts && python gen_xlsx_comp.py\n"
       "필요 패키지: openpyxl. 출력: ~/Documents/hr-samples/나린테크 급여근태 대장(2026H1).xlsx")
    kv("결정론",
       "roster.SEED = %d 고정 + 연차 사용일수의 사번 SHA-256 시드 고정. 두 축 모두 프로세스 "
       "환경에 의존하지 않으므로 산출물은 바이트 단위 데이터가 매 실행 동일하다." % R.SEED)
    kv("자체 검증",
       "스크립트 말미의 verify() 블록이 행 수(212/212), 시트②의 인원 합계 212·밴드이탈 0, "
       "시트③의 인원 합계 = data.half(2026,1)[\"n\"], 가중평균 인상률의 data.py 대조(소수 2자리), "
       "잔여 연차 음수 0건을 검사하고 하나라도 어긋나면 예외를 던진다.")

    for r in range(1, ws.max_row + 1):
        ws.row_dimensions[r].height = None
    return ws


# ══════════════════════════════════════════════════════════════════
# § 검증
# ══════════════════════════════════════════════════════════════════
def verify(recs, hf, w_raise, w_bonus):
    """요구된 검증 항목을 전부 실행하고 (통과여부, 리포트행) 을 반환한다."""
    checks = []

    def chk(name, ok, detail):
        checks.append((name, ok, detail))

    n_pay = len(recs)
    chk("급여대장 행 수 == 212", n_pay == 212, "%d행" % n_pay)

    n_leave = len(recs)   # 연차현황도 동일 레코드셋
    chk("연차현황 행 수 == 212", n_leave == 212, "%d행" % n_leave)

    band_sum = sum(len([r for r in recs if r["level"] == lvl]) for lvl in LEVEL_ORDER)
    chk("직급별급여분포 인원 합계 == 212", band_sum == 212, "%d명" % band_sum)

    out = sum(1 for r in recs
              if r["annual"] < LEVEL_BAND[r["level"]][0] or r["annual"] > LEVEL_BAND[r["level"]][1])
    chk("밴드이탈인원 == 0", out == 0, "%d명" % out)

    graded = [r for r in recs if r["grade"]]
    chk("등급별처우 인원 합계 == half(2026,1)['n']", len(graded) == hf["n"],
        "%d == %d" % (len(graded), hf["n"]))

    ref_raise = D.weighted_raise(hf)
    ok_raise = round(w_raise, 2) == round(ref_raise, 2)
    chk("가중평균 인상률 == data.weighted_raise (소수 2자리)", ok_raise,
        "시트3 %.2f%% vs data.py %.2f%%" % (w_raise, ref_raise))

    ref_bonus = D.weighted_bonus(hf)
    ok_bonus = round(w_bonus, 2) == round(ref_bonus, 2)
    chk("가중평균 상여율 == data.weighted_bonus (소수 2자리)", ok_bonus,
        "시트3 %.2f%% vs data.py %.2f%%" % (w_bonus, ref_bonus))

    neg = sum(1 for r in recs if r["remain"] < 0)
    chk("연차 잔여일수 음수 행 == 0", neg == 0, "%d행" % neg)

    # 보조 검증 — 등급 분포가 data.HALVES와 정확히 일치하는지
    dist_ok = all(len([r for r in graded if r["grade"] == g]) == hf["dist"][g] for g in D.GRADES)
    chk("등급별 인원 == data.HALVES 2026-H1 분포", dist_ok,
        " ".join("%s%d" % (g, len([r for r in graded if r["grade"] == g])) for g in D.GRADES))

    return checks


# ══════════════════════════════════════════════════════════════════
def main():
    print("── 로스터 생성 (roster.build, seed=%d) ──" % R.SEED)
    people, active = R.build()
    print("전체 %d명 / 재직 %d명 / 퇴사 %d명"
          % (len(people), len(active), len(people) - len(active)))

    hf = D.half(*HALF_KEY)
    recs = build_records(active)
    n_excluded = sum(1 for r in recs if not r["grade"])
    print("2026H1 평가대상 %d명 / 평가제외(수습) %d명" % (hf["n"], n_excluded))

    wb = Workbook()
    wb.remove(wb.active)          # 기본 시트 제거

    sheet_payroll(wb, recs)
    sheet_band(wb, recs)
    _ws3, w_raise, w_bonus = sheet_grade(wb, recs, hf)
    sheet_leave(wb, recs)
    sheet_dept_leave(wb, recs)
    _ws6, avg_monthly = sheet_trend(wb, recs)
    sheet_readme(wb, recs, hf, avg_monthly, n_excluded)

    # ── 검증 ──
    print()
    print("── 검증 ──")
    checks = verify(recs, hf, w_raise, w_bonus)
    for name, ok, detail in checks:
        print("  [%s] %-46s %s" % ("PASS" if ok else "FAIL", name, detail))
    failed = [c for c in checks if not c[1]]
    if failed:
        raise SystemExit("검증 실패 %d건 — 워크북을 저장하지 않는다." % len(failed))

    os.makedirs(OUT_DIR, exist_ok=True)
    wb.save(OUT_PATH)
    size = os.path.getsize(OUT_PATH)
    print()
    print("저장: %s (%.1f KB)" % (OUT_PATH, size / 1024.0))
    print("시트별 행 수:")
    for ws in wb.worksheets:
        print("  %-14s %d행 × %d열" % (ws.title, ws.max_row, ws.max_column))


if __name__ == "__main__":
    main()
