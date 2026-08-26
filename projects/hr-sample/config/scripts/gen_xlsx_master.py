# -*- coding: utf-8 -*-
"""(주)나린테크 「인사마스터 대장(2026H1)」 XLSX 생성기.

실제 인사팀이 유지하는 마스터 대장을 흉내 낸 워크북 1종(시트 6개)을 만든다.
모든 수치는 `roster.py`(개인 단위 로스터)와 `data.py`(집계 단일 진실원)에서만
파생된다 — 이 스크립트는 값을 새로 지어내지 않고 배치·서식만 담당한다.

시트 구성
  1. 재직자마스터   — 재직자 212명, 1인 1행 (반기 등급 5개 차수 가로 전개)
  2. 퇴사자마스터   — 퇴사자 72명, 1인 1행
  3. 평가이력       — 반기 차수 × 개인 롱 포맷 888행
  4. 부서별집계     — 본부별 요약 + 합계 행
  5. 등급분포검증   — 차수 × 등급 25행 (규정 QUOTA 대비 실제)
  6. README         — 시트 설명·출처·재현 방법·가상 데이터 고지

실행:
    cd <이 파일이 있는 디렉터리>
    <venv>/bin/python gen_xlsx_master.py
"""
import os
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo  # noqa: F401  (확장 여지)

import data as D
import roster as R

# ══════════════════════════════════════════════════════════════════
# § 상수 — 출력 경로·기준일·서식 팔레트
# ══════════════════════════════════════════════════════════════════

OUT_DIR = os.path.expanduser("~/Documents/hr-samples")
OUT_NAME = "나린테크 인사마스터 대장(2026H1).xlsx"
OUT_PATH = os.path.join(OUT_DIR, OUT_NAME)

# 대장의 기준일 — 2026 상반기 말. 근속·직급·연봉은 모두 이 시점 기준이다.
ASOF = date(2026, 6, 30)

# 반기 차수 키 (연, 반기) — data.HALVES에서 그대로 가져온다.
HALF_KEYS = [(row[0], row[1]) for row in D.HALVES]
HALF_LABELS = ["%dH%d" % (y, h) for y, h in HALF_KEYS]

# 반기 종료일 — 평가 시점 근속·직급 계산의 기준
HALF_END = {(y, h): (date(y, 6, 30) if h == 1 else date(y, 12, 31)) for y, h in HALF_KEYS}

# 숫자 표시 서식
FMT_DATE = "YYYY-MM-DD"
FMT_MONEY = "#,##0"          # 금액(만원) 천단위 구분
FMT_PCT = '0.0"%"'           # 백분율 — 값은 8.0 처럼 이미 % 단위 숫자
FMT_PCT0 = '0"%"'            # 정수 백분율(규정 비율)
FMT_DIFF = '+0.0"%p";-0.0"%p";0.0"%p"'   # 편차 — 부호 항상 표기
FMT_YEARS = "0.0"

# 색 팔레트
C_HEADER = "305496"          # 헤더 배경 (진한 남색)
C_HEADER_FG = "FFFFFF"
C_TOTAL = "D9E1F2"           # 합계 행 배경 (연한 파랑)
C_ZEBRA = "F7F9FC"           # 홀수 행 옅은 줄무늬

# 등급별 채우기 — S는 진한 파랑(흰 글씨), D는 붉은 계열
GRADE_FILL = {
    "S": ("1F4E79", "FFFFFF", True),
    "A": ("9DC3E6", "1F3864", False),
    "B": ("EDEDED", "3F3F3F", False),
    "C": ("FFE699", "7F6000", False),
    "D": ("F4777A", "5C0B0E", True),
}

THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

LEFT = Alignment(horizontal="left", vertical="center")
CENTER = Alignment(horizontal="center", vertical="center")
RIGHT = Alignment(horizontal="right", vertical="center")


# ══════════════════════════════════════════════════════════════════
# § 공통 헬퍼
# ══════════════════════════════════════════════════════════════════

def level_label(p, asof=ASOF):
    """직급을 'L3(책임)' 형태로 — 코드만으로는 대장에서 읽히지 않는다."""
    code = R.level_of(p, asof)
    name = next(n for c, n, *_ in D.GRADE_LEVELS if c == code)
    return "%s(%s)" % (code, name)


def latest_grade(p):
    """가장 최근 차수의 등급. 평가 이력이 없으면 None(수습 중 신규 입사자)."""
    if not p["grades"]:
        return None
    return p["grades"][sorted(p["grades"])[-1]]


def tenure_at(p, asof):
    """임의 시점 기준 근속연수(년, 소수 1자리)."""
    return round((asof - p["hire_date"]).days / 365.25, 1)


def write_header(ws, headers, row=1):
    """헤더 행 — 굵게 + 배경색 + 가운데 정렬 + 테두리."""
    fill = PatternFill("solid", fgColor=C_HEADER)
    font = Font(bold=True, color=C_HEADER_FG, size=10)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=row, column=i, value=h)
        c.fill = fill
        c.font = font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[row].height = 26


def set_widths(ws, widths):
    """열 너비 일괄 지정."""
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def finish_grid(ws, n_rows, n_cols, header_row=1, freeze="A2", autofilter=True):
    """틀고정 + 자동필터 + 헤더 반복 인쇄 설정을 한 번에."""
    ws.freeze_panes = freeze
    if autofilter and n_rows > 0:
        ws.auto_filter.ref = "A%d:%s%d" % (
            header_row, get_column_letter(n_cols), header_row + n_rows)
    ws.print_title_rows = "%d:%d" % (header_row, header_row)


def paint_grade(cell, grade):
    """등급 셀에 색을 입힌다 (조건부 서식 대신 직접 채우기 — 값이 고정이라 결정론적)."""
    if not grade:
        return
    bg, fg, bold = GRADE_FILL[grade]
    cell.fill = PatternFill("solid", fgColor=bg)
    cell.font = Font(color=fg, bold=bold, size=10)
    cell.alignment = CENTER


def style_row(ws, r, n_cols, aligns, zebra=False):
    """행 단위 정렬·테두리·줄무늬 적용. aligns는 열별 Alignment 리스트."""
    fill = PatternFill("solid", fgColor=C_ZEBRA) if zebra else None
    for i in range(1, n_cols + 1):
        c = ws.cell(row=r, column=i)
        c.border = BORDER
        # 정렬이 아직 안 잡힌 셀에만 열 규칙을 적용한다
        if not c.alignment or c.alignment.horizontal is None:
            c.alignment = aligns[i - 1]
        if fill is not None and (c.fill is None or c.fill.fill_type is None):
            c.fill = fill


# ══════════════════════════════════════════════════════════════════
# § 시트 1 — 재직자마스터
# ══════════════════════════════════════════════════════════════════

S1_HEADERS = (["사번", "본부", "직급", "입사일", "근속연수", "채용채널", "수습통과"]
              + HALF_LABELS
              + ["최근등급", "연봉(만원)", "조직이동이력"])
S1_WIDTHS = [9, 15, 12, 12, 10, 11, 9] + [9] * 5 + [10, 12, 24]


def sheet_active(wb, active):
    """재직자 전원 1인 1행. 등급 열은 해당 차수 평가 대상이 아니었으면 빈칸."""
    ws = wb.create_sheet("재직자마스터")
    write_header(ws, S1_HEADERS)
    n_cols = len(S1_HEADERS)

    # 정렬 규칙: 텍스트 좌측, 숫자 우측, 등급·플래그는 가운데
    aligns = ([LEFT, LEFT, LEFT, CENTER, RIGHT, LEFT, CENTER]
              + [CENTER] * 5 + [CENTER, RIGHT, LEFT])

    r = 1
    for p in sorted(active, key=lambda x: x["emp_id"]):
        r += 1
        move = ("%s → %s" % (p["prev_dept"], p["dept"])) if p.get("dept_changed") else ""
        vals = [p["emp_id"], p["dept"], level_label(p), p["hire_date"],
                R.tenure_years(p, ASOF), p["channel"],
                "Y" if p["probation_passed"] else "N"]
        for k in HALF_KEYS:
            vals.append(p["grades"].get(k, ""))
        vals += [latest_grade(p) or "", R.salary(p, ASOF), move]

        for i, v in enumerate(vals, start=1):
            ws.cell(row=r, column=i, value=v)

        ws.cell(row=r, column=4).number_format = FMT_DATE
        ws.cell(row=r, column=5).number_format = FMT_YEARS
        ws.cell(row=r, column=n_cols - 1).number_format = FMT_MONEY

        style_row(ws, r, n_cols, aligns, zebra=(r % 2 == 1))
        # 등급 열(8~12) + 최근등급 열에 색 입히기
        for i, k in enumerate(HALF_KEYS):
            paint_grade(ws.cell(row=r, column=8 + i), p["grades"].get(k))
        paint_grade(ws.cell(row=r, column=13), latest_grade(p))

    set_widths(ws, S1_WIDTHS)
    finish_grid(ws, r - 1, n_cols)
    return r - 1


# ══════════════════════════════════════════════════════════════════
# § 시트 2 — 퇴사자마스터
# ══════════════════════════════════════════════════════════════════

S2_HEADERS = ["사번", "본부", "직급", "입사일", "퇴사일", "근속연수", "채용채널",
              "퇴사구분(자발/비자발)", "퇴사사유", "최종등급", "퇴사분기"]
S2_WIDTHS = [9, 15, 12, 12, 12, 10, 11, 18, 16, 10, 11]


def sheet_leavers(wb, leavers):
    """퇴사자 전원. 근속연수·직급은 퇴사일 시점 기준(tenure_years가 term_date 우선)."""
    ws = wb.create_sheet("퇴사자마스터")
    write_header(ws, S2_HEADERS)
    n_cols = len(S2_HEADERS)
    aligns = [LEFT, LEFT, LEFT, CENTER, CENTER, RIGHT, LEFT, CENTER, LEFT, CENTER, CENTER]

    r = 1
    for p in sorted(leavers, key=lambda x: (x["term_date"], x["emp_id"])):
        r += 1
        y, q = p["term_quarter"]
        # 퇴사자는 퇴사일 시점으로 직급을 확정한다 — 기준일(ASOF)을 쓰면 이미 나간
        # 사람의 근속이 계속 늘어난다.
        vals = [p["emp_id"], p["dept"], level_label(p, p["term_date"]),
                p["hire_date"], p["term_date"], R.tenure_years(p),
                p["channel"], p["term_type"], p["term_reason"],
                latest_grade(p) or "", "%d-Q%d" % (y, q)]
        for i, v in enumerate(vals, start=1):
            ws.cell(row=r, column=i, value=v)

        ws.cell(row=r, column=4).number_format = FMT_DATE
        ws.cell(row=r, column=5).number_format = FMT_DATE
        ws.cell(row=r, column=6).number_format = FMT_YEARS

        style_row(ws, r, n_cols, aligns, zebra=(r % 2 == 1))
        paint_grade(ws.cell(row=r, column=10), latest_grade(p))
        # 비자발 퇴사는 구분 열을 옅게 강조 — 대장에서 눈으로 골라내기 쉽게
        if p["term_type"] == "비자발":
            c = ws.cell(row=r, column=8)
            c.fill = PatternFill("solid", fgColor="FCE4D6")
            c.font = Font(color="843C0C", size=10)

    set_widths(ws, S2_WIDTHS)
    finish_grid(ws, r - 1, n_cols)
    return r - 1


# ══════════════════════════════════════════════════════════════════
# § 시트 3 — 평가이력 (롱 포맷)
# ══════════════════════════════════════════════════════════════════

S3_HEADERS = ["사번", "평가차수", "본부", "직급", "등급", "등급점수",
              "근속연수(평가시점)", "인상률(%)", "상여지급률(%)"]
S3_WIDTHS = [9, 11, 15, 12, 8, 10, 17, 11, 13]


def sheet_reviews(wb, people):
    """(개인 × 평가차수) 롱 포맷. 5개 차수 전부 포함 — 합계 888행이어야 한다.

    본부는 현재 소속(재직자는 2026H1 기준, 퇴사자는 퇴사 시점)이다. roster.py의
    부서 리밸런싱이 등급 배정 뒤에 일어나므로 과거 차수의 소속 이력은 남아 있지
    않다 — 이동 이력은 시트 1의 `조직이동이력` 열로만 추적한다.
    """
    ws = wb.create_sheet("평가이력")
    write_header(ws, S3_HEADERS)
    n_cols = len(S3_HEADERS)
    aligns = [LEFT, CENTER, LEFT, LEFT, CENTER, RIGHT, RIGHT, RIGHT, RIGHT]

    rows = []
    for p in people:
        for k, g in p["grades"].items():
            rows.append((k, p, g))
    # 차수 → 사번 순으로 정렬해 사람 눈으로도 훑을 수 있게 한다
    rows.sort(key=lambda t: (t[0], t[1]["emp_id"]))

    per_half = {}
    r = 1
    for k, p, g in rows:
        r += 1
        per_half[k] = per_half.get(k, 0) + 1
        asof = HALF_END[k]
        vals = [p["emp_id"], "%dH%d" % k, p["dept"], level_label(p, asof), g,
                R.GRADE_SCORE[g], tenure_at(p, asof),
                D.GRADE_TO_RAISE[g], float(D.GRADE_TO_BONUS[g])]
        for i, v in enumerate(vals, start=1):
            ws.cell(row=r, column=i, value=v)

        ws.cell(row=r, column=7).number_format = FMT_YEARS
        ws.cell(row=r, column=8).number_format = FMT_PCT
        ws.cell(row=r, column=9).number_format = FMT_PCT

        style_row(ws, r, n_cols, aligns, zebra=(r % 2 == 1))
        paint_grade(ws.cell(row=r, column=5), g)

    set_widths(ws, S3_WIDTHS)
    finish_grid(ws, r - 1, n_cols)
    return r - 1, per_half


# ══════════════════════════════════════════════════════════════════
# § 시트 4 — 부서별집계
# ══════════════════════════════════════════════════════════════════

S4_HEADERS = (["본부", "재직인원", "평균근속", "평균연봉(만원)"] + D.GRADES
              + ["평균등급점수", "2026H1 자발퇴사"])
S4_WIDTHS = [15, 10, 10, 14, 6, 6, 6, 6, 6, 12, 15]


def sheet_dept(wb, active, leavers):
    """본부별 재직 요약 + 합계 행.

    - S~D 열: 재직자의 **최근등급** 분포. 2026-Q2 입사자 16명은 아직 평가 대상이
      아니어서 등급이 없으므로 S~D 합(196)은 재직인원(212)과 다르다 — 의도된 차이다.
    - 2026H1 자발퇴사: 2026-Q1·Q2에 자발 퇴사한 인원(합계 11명).
    """
    ws = wb.create_sheet("부서별집계")
    write_header(ws, S4_HEADERS)
    n_cols = len(S4_HEADERS)
    aligns = [LEFT] + [RIGHT] * (n_cols - 1)

    # 2026 상반기(Q1·Q2) 자발 퇴사자를 본부별로 집계
    vol_2026h1 = {}
    for p in leavers:
        if p["term_quarter"] in ((2026, 1), (2026, 2)) and p["term_type"] == "자발":
            vol_2026h1[p["dept"]] = vol_2026h1.get(p["dept"], 0) + 1

    r = 1
    tot_n = 0
    tot_ten = 0.0
    tot_sal = 0
    tot_dist = {g: 0 for g in D.GRADES}
    tot_score_sum = 0
    tot_score_n = 0
    tot_vol = 0

    for dept, _ in D.DEPT_SHARE:
        grp = [p for p in active if p["dept"] == dept]
        r += 1
        dist = {g: 0 for g in D.GRADES}
        score_sum, score_n = 0, 0
        for p in grp:
            g = latest_grade(p)
            if g:
                dist[g] += 1
                score_sum += R.GRADE_SCORE[g]
                score_n += 1
        avg_ten = sum(R.tenure_years(p, ASOF) for p in grp) / len(grp)
        avg_sal = sum(R.salary(p, ASOF) for p in grp) / len(grp)
        avg_score = score_sum / score_n if score_n else 0.0
        vol = vol_2026h1.get(dept, 0)

        vals = ([dept, len(grp), round(avg_ten, 1), round(avg_sal)]
                + [dist[g] for g in D.GRADES] + [round(avg_score, 2), vol])
        for i, v in enumerate(vals, start=1):
            ws.cell(row=r, column=i, value=v)
        ws.cell(row=r, column=3).number_format = FMT_YEARS
        ws.cell(row=r, column=4).number_format = FMT_MONEY
        ws.cell(row=r, column=10).number_format = "0.00"
        style_row(ws, r, n_cols, aligns, zebra=(r % 2 == 1))
        for i, g in enumerate(D.GRADES):
            paint_grade(ws.cell(row=r, column=5 + i), g)
            ws.cell(row=r, column=5 + i).alignment = CENTER

        tot_n += len(grp)
        tot_ten += sum(R.tenure_years(p, ASOF) for p in grp)
        tot_sal += sum(R.salary(p, ASOF) for p in grp)
        for g in D.GRADES:
            tot_dist[g] += dist[g]
        tot_score_sum += score_sum
        tot_score_n += score_n
        tot_vol += vol

    # ── 합계 행 — 재직인원 합계는 반드시 212 ──
    r += 1
    vals = (["합계", tot_n, round(tot_ten / tot_n, 1), round(tot_sal / tot_n)]
            + [tot_dist[g] for g in D.GRADES]
            + [round(tot_score_sum / tot_score_n, 2), tot_vol])
    for i, v in enumerate(vals, start=1):
        c = ws.cell(row=r, column=i, value=v)
        c.fill = PatternFill("solid", fgColor=C_TOTAL)
        c.font = Font(bold=True, size=10)
        c.border = BORDER
        c.alignment = aligns[i - 1]
    ws.cell(row=r, column=3).number_format = FMT_YEARS
    ws.cell(row=r, column=4).number_format = FMT_MONEY
    ws.cell(row=r, column=10).number_format = "0.00"

    set_widths(ws, S4_WIDTHS)
    ws.freeze_panes = "B2"
    return tot_n, r


# ══════════════════════════════════════════════════════════════════
# § 시트 5 — 등급분포검증
# ══════════════════════════════════════════════════════════════════

S5_HEADERS = ["평가차수", "등급", "규정비율(%)", "실제인원", "실제비율(%)", "편차(%p)"]
S5_WIDTHS = [11, 8, 13, 11, 12, 12]


def sheet_quota(wb):
    """차수 × 등급 25행 — 「인사평가 운영지침」 강제 분포(QUOTA) 대비 실제치."""
    ws = wb.create_sheet("등급분포검증")
    write_header(ws, S5_HEADERS)
    n_cols = len(S5_HEADERS)
    aligns = [CENTER, CENTER, RIGHT, RIGHT, RIGHT, RIGHT]

    r = 1
    per_half = {}
    for key in HALF_KEYS:
        hf = D.half(*key)
        for g in D.GRADES:
            r += 1
            cnt = hf["dist"][g]
            per_half[key] = per_half.get(key, 0) + cnt
            act = cnt / hf["n"] * 100
            quota = D.QUOTA[g] * 100
            vals = ["%dH%d" % key, g, quota, cnt, round(act, 1), round(act - quota, 1)]
            for i, v in enumerate(vals, start=1):
                ws.cell(row=r, column=i, value=v)
            ws.cell(row=r, column=3).number_format = FMT_PCT0
            ws.cell(row=r, column=5).number_format = FMT_PCT
            ws.cell(row=r, column=6).number_format = FMT_DIFF
            style_row(ws, r, n_cols, aligns, zebra=(r % 2 == 1))
            paint_grade(ws.cell(row=r, column=2), g)
        # 차수 경계에 굵은 아래 테두리 — 5행 묶음이 눈에 들어오게
        for i in range(1, n_cols + 1):
            c = ws.cell(row=r, column=i)
            c.border = Border(left=THIN, right=THIN, top=THIN,
                              bottom=Side(style="medium", color="808080"))

    set_widths(ws, S5_WIDTHS)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(n_cols), r)
    return r - 1, per_half


# ══════════════════════════════════════════════════════════════════
# § 시트 6 — README
# ══════════════════════════════════════════════════════════════════

def sheet_readme(wb, counts):
    """워크북 사용 설명서 — 시트별 정의, 데이터 출처, 재현 방법, 고지."""
    ws = wb.create_sheet("README")
    write_header(ws, ["항목", "설명"])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110

    rows = [
        ("워크북명", "%s 인사마스터 대장(2026H1)" % D.COMPANY),
        ("기준일", "2026-06-30 (2026 상반기 말). 근속연수·직급·연봉은 모두 이 시점 기준이다."),
        ("생성 스크립트", "projects/hr-sample/config/scripts/gen_xlsx_master.py"),
        ("데이터 출처",
         "roster.py (개인 단위 로스터, SEED=%d) + data.py (집계 단일 진실원). "
         "급여 밴드·인상률·상여 지급률은 「나린테크 급여규정(2026 개정)」 제3조·제8조·제9조에 대응한다."
         % R.SEED),
        ("재현 방법",
         "cd projects/hr-sample/config/scripts && <venv>/bin/python gen_xlsx_master.py  → "
         "~/Documents/hr-samples/ 아래에 동일 파일이 다시 생성된다. "
         "roster.py의 SEED가 고정되어 있어 몇 번을 돌려도 같은 결과가 나온다(결정론적)."),
        ("", ""),
        ("시트 1 재직자마스터",
         "2026-06-30 기준 재직자 %d명 전원, 1인 1행. 2024H1~2026H1 5개 평가 차수의 등급을 "
         "가로로 전개한다. 해당 차수에 평가 대상이 아니었던 경우(입사 전이거나 당 분기 입사로 "
         "수습 중)는 빈칸이다. `조직이동이력`은 부서 재배치로 소속이 바뀐 인원에만 "
         "'이전본부 → 현재본부' 형태로 채워진다." % counts["active"]),
        ("시트 2 퇴사자마스터",
         "2024-Q1~2026-Q2에 퇴사한 %d명 전원. 근속연수·직급은 퇴사일 시점 기준으로 확정한다. "
         "`퇴사구분`은 자발/비자발, `최종등급`은 재직 중 마지막으로 받은 평가 등급이다. "
         "수습 미전환자는 「취업규칙」 제12조의 수습 3개월을 반영해 입사 다음 분기에 퇴사 처리된다."
         % counts["leavers"]),
        ("시트 3 평가이력",
         "(개인 × 평가차수) 롱 포맷 %d행. 차수별 평가 대상은 %s이며 "
         "합이 %d행이 된다. `인상률(%%)`·`상여지급률(%%)`은 data.py의 GRADE_TO_RAISE/"
         "GRADE_TO_BONUS에서 그대로 가져온 제도 연동값이다(급여규정 제8조·제9조). "
         "`본부`는 현재 소속 기준 — 과거 차수 시점의 소속 이력은 보존하지 않는다."
         % (counts["reviews"],
            " / ".join("%s %d명" % (lb, D.half(*k)["n"]) for lb, k in zip(HALF_LABELS, HALF_KEYS)),
            counts["reviews"])),
        ("시트 4 부서별집계",
         "본부별 재직 요약 + 합계 행. `재직인원` 합계는 %d으로 시트 1의 행 수와 일치한다. "
         "S~D 열은 재직자의 **최근등급** 분포이며, 2026-Q2 입사자 %d명은 아직 평가 대상이 "
         "아니어서 등급이 없다 — 따라서 S~D 합(%d)은 재직인원(%d)보다 작다. `평균등급점수`는 "
         "S=4/A=3/B=2/C=1/D=0으로 환산한 값이고, 등급 보유자만 분모에 넣는다."
         % (counts["dept_total"], counts["ungraded"],
            counts["dept_total"] - counts["ungraded"], counts["dept_total"])),
        ("시트 5 등급분포검증",
         "5개 차수 × 5등급 = %d행. `규정비율`은 「인사평가 운영지침」의 강제 분포 "
         "(S10/A20/B50/C15/D5, data.QUOTA), `실제인원`은 data.HALVES의 차수별 실제 배정 "
         "인원이다. `편차(%%p)`가 0이 아닌 것은 정수 인원을 비율에 정확히 맞출 수 없기 "
         "때문이며, 절댓값이 큰 항목이 운영지침 준수 점검의 착안점이 된다." % counts["quota"]),
        ("시트 6 README", "지금 보고 있는 시트."),
        ("", ""),
        ("자기 검증", "생성 시 다음을 강제한다 — %s" % counts["verify_line"]),
        ("서식 규칙",
         "날짜 YYYY-MM-DD, 금액 천단위 구분(#,##0, 단위: 만원), 백분율 0.0\"%\". "
         "숫자 열은 우측, 텍스트 열은 좌측 정렬. 시트 1·2·3은 헤더 틀고정 + 자동필터. "
         "평가 등급 셀은 S(진한 파랑) → D(붉은색)로 색을 입혀 분포를 눈으로 훑을 수 있게 했다."),
        ("", ""),
        ("고지", D.FIXTURE_NOTE + " 사번·부서·평가등급·연봉을 포함한 모든 레코드는 "
                 "결정론적 시뮬레이션으로 생성한 가상 목업 데이터이며, 실존 기업·인물·조직과 "
                 "무관하다. 실제 인사 의사결정의 근거로 사용해서는 안 된다."),
    ]

    r = 1
    for k, v in rows:
        r += 1
        a = ws.cell(row=r, column=1, value=k)
        b = ws.cell(row=r, column=2, value=v)
        a.font = Font(bold=True, size=10)
        a.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        b.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
        if k.startswith("시트"):
            a.fill = PatternFill("solid", fgColor="E2EFDA")
        if k == "고지":
            a.fill = PatternFill("solid", fgColor="FCE4D6")
            b.font = Font(italic=True, size=10)
        if k:
            a.border = BORDER
            b.border = BORDER
        # 긴 설명은 행 높이를 넉넉히
        ws.row_dimensions[r].height = max(18, 15 * (len(v) // 95 + 1))

    ws.freeze_panes = "A2"
    return r - 1


# ══════════════════════════════════════════════════════════════════
# § 메인 — 생성 + 검증
# ══════════════════════════════════════════════════════════════════

def main():
    # 1) 집계 모델 자기 검증 (data.py) — 여기서 깨지면 아래는 볼 필요도 없다
    base_errs = D.verify() + D.verify_ext()
    if base_errs:
        print("data.py 검증 실패:")
        for e in base_errs:
            print("  -", e)
        raise SystemExit(1)

    # 2) 로스터 생성 + 로스터 자기 검증
    people, active = R.build()
    r_errs, _, _ = R.verify(people, active)
    if r_errs:
        print("roster.py 검증 실패:")
        for e in r_errs[:10]:
            print("  -", e)
        raise SystemExit(1)

    leavers = [p for p in people if p["term_date"] is not None]
    ungraded = sum(1 for p in active if not p["grades"])

    # 3) 워크북 조립
    wb = Workbook()
    wb.remove(wb.active)   # 기본 시트 제거

    n1 = sheet_active(wb, active)
    n2 = sheet_leavers(wb, leavers)
    n3, rev_per_half = sheet_reviews(wb, people)
    dept_total, _ = sheet_dept(wb, active, leavers)
    n5, quota_per_half = sheet_quota(wb)

    verify_line = ("재직자마스터 212행 / 퇴사자마스터 72행 / 평가이력 888행 / "
                   "부서별집계 재직인원 합계 212 / 등급분포검증 차수별 실제인원 합 = "
                   "data.HALVES의 n")
    sheet_readme(wb, {
        "active": n1, "leavers": n2, "reviews": n3, "dept_total": dept_total,
        "quota": n5, "ungraded": ungraded, "verify_line": verify_line,
    })

    wb.properties.title = "나린테크 인사마스터 대장(2026H1)"
    wb.properties.creator = "인사팀 (가상 목업)"
    wb.properties.description = D.FIXTURE_NOTE

    os.makedirs(OUT_DIR, exist_ok=True)
    wb.save(OUT_PATH)

    # ── 4) 검증 블록 — 하나라도 어긋나면 실패로 종료 ──────────────
    checks = []
    checks.append(("재직자마스터 행 수", n1, 212))
    checks.append(("퇴사자마스터 행 수", n2, 72))
    expected_rev = sum(D.half(*k)["n"] for k in HALF_KEYS)
    checks.append(("평가이력 행 수", n3, expected_rev))
    checks.append(("부서별집계 재직인원 합계", dept_total, 212))
    checks.append(("등급분포검증 행 수", n5, len(HALF_KEYS) * len(D.GRADES)))
    for k in HALF_KEYS:
        n = D.half(*k)["n"]
        checks.append(("등급분포검증 %dH%d 실제인원 합" % k, quota_per_half[k], n))
        checks.append(("평가이력 %dH%d 행 수" % k, rev_per_half.get(k, 0), n))

    print("생성 완료: %s (%.1f KB)" % (OUT_PATH, os.path.getsize(OUT_PATH) / 1024.0))
    print("시트: %s" % ", ".join(wb.sheetnames))
    print()
    print("── 검증 ──────────────────────────────────────────")
    failed = 0
    for label, got, want in checks:
        ok = (got == want)
        if not ok:
            failed += 1
        print("  [%s] %-32s %5s (기대 %s)" % ("OK" if ok else "FAIL", label, got, want))
    print("──────────────────────────────────────────────────")
    if failed:
        print("검증 실패 %d건 — 산출물을 신뢰할 수 없다." % failed)
        raise SystemExit(1)
    print("전체 %d건 통과." % len(checks))


if __name__ == "__main__":
    main()
