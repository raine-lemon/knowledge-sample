#!/usr/bin/env python3
# /// script
# dependencies = ["python-pptx"]
# ///
"""P1: python-pptx PPTX -> MD 직행 변환 (pptx2md-ingest).

usage:
  uv run p1-convert.py <file.pptx>           # 변환된 MD를 stdout으로
  uv run p1-convert.py <file.pptx> --stats   # 측정 통계만 JSON으로 (§1 전략 판정용)

(uv가 PEP 723 메타데이터로 python-pptx를 자동 준비한다. 이미 설치된 환경이면
 python3로 직접 실행해도 된다.)

MD 산출 규칙 (SKILL.md §2):
  - 첫 줄은 원제목 H1 — 1번 슬라이드의 첫 텍스트 줄, 없으면 파일명
  - 슬라이드마다 `<!-- slide N -->` 마커
  - 제목 도형은 `## `, 부제는 `_기울임_`
  - 불릿은 `-` / `  -` (단락 level + 불릿 글리프 중 깊은 쪽)
  - 표 도형은 md 표, 발표자 노트는 `> 발표자 노트:` 인용 블록
  - 그룹 도형은 재귀 순회, SmartArt는 mc:AlternateContent 폴백에서 텍스트 수집
  - 전사 불가한 이미지/차트는 `[IMAGE: 이름]` / `[CHART: 이름]` 마커로 남기고 센다

--stats JSON 키:
  slides, notes, images, tables, charts, chars(공백 제외), chars_per_slide,
  sparse_slides(공백 제외 50자 미만 슬라이드 수), sparse_slide_numbers(P3 대상)
"""
import json
import sys
from pathlib import Path

SPARSE_CHARS = 50  # 슬라이드당 공백 제외 문자수 — 이 미만이면 "텍스트 희소" (§1)

BULLET_L0 = "•●▪■◆‧·*"          # 1단 불릿 글리프
BULLET_L1 = "–—-‣◦○▫"           # 2단 불릿 글리프
MC_ALT = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"


def strip_bullet(s):
    """선두 불릿 글리프를 떼고 (본문, 글리프 단계) 반환. 글리프 없으면 (본문, None)."""
    if len(s) >= 2 and s[1] in " \t\u00a0":
        if s[0] in BULLET_L0:
            return s[2:].strip(), 0
        if s[0] in BULLET_L1:
            return s[2:].strip(), 1
    return s, None


def para_lines(text_frame):
    """텍스트 프레임 -> [(level, text, bulleted)]. 줄바꿈(a:br)은 같은 level로 분해."""
    out = []
    for para in text_frame.paragraphs:
        raw = (para.text or "").replace("\x0b", "\n")
        for piece in raw.split("\n"):
            piece = piece.strip()
            if not piece:
                continue
            body, rank = strip_bullet(piece)
            if not body:
                continue
            level = para.level or 0
            if rank is not None and level == 0:
                level = rank
            out.append((level, body, rank is not None or (para.level or 0) > 0))
    return out


def a_texts(element):
    """요소 하위 모든 <a:t>를 순서 유지 dedupe로 수집 (SmartArt 폴백 등)."""
    seen, out = set(), []
    for node in element.iter("{http://schemas.openxmlformats.org/drawingml/2006/main}t"):
        val = (node.text or "").strip()
        if val and val not in seen:
            seen.add(val)
            out.append(val)
    return out


def walk(shapes):
    """(kind, shape) 를 문서 순서대로. 그룹 도형은 재귀적으로 펼친다."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    for shape in shapes:
        stype = shape.shape_type
        if stype == MSO_SHAPE_TYPE.GROUP:
            yield from walk(shape.shapes)
        elif getattr(shape, "has_table", False):
            yield "table", shape
        elif getattr(shape, "has_chart", False):
            yield "chart", shape
        elif shape.has_text_frame:
            if shape.text_frame.text.strip():
                yield "text", shape
        elif stype in (MSO_SHAPE_TYPE.PICTURE, MSO_SHAPE_TYPE.LINKED_PICTURE,
                       MSO_SHAPE_TYPE.MEDIA):
            yield "image", shape
        elif a_texts(shape._element):
            yield "xmltext", shape          # SmartArt·기타 graphicFrame 안의 텍스트
        elif shape.__class__.__name__ == "GraphicFrame":
            yield "image", shape            # 전사 불가한 그래픽 프레임


def table_chunk(shape):
    rows = [[(cell.text or "").strip().replace("|", "\\|").replace("\n", "<br>")
             for cell in row.cells] for row in shape.table.rows]
    if not rows:
        return None
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |",
             "|" + "|".join(["---"] * width) + "|"]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


def render_lines(lines):
    """[(level, text, bulleted)] -> 청크 리스트. 연속 불릿은 한 청크로 묶는다."""
    chunks, bullets = [], []
    for level, text, bulleted in lines:
        if bulleted:
            bullets.append("  " * level + "- " + text)
        else:
            if bullets:
                chunks.append("\n".join(bullets))
                bullets = []
            chunks.append(text)
    if bullets:
        chunks.append("\n".join(bullets))
    return chunks


def merge_bullets(chunks):
    """도형이 갈려 끊긴 인접 불릿 청크를 하나로 합친다 (그룹 도형 레이아웃 대응)."""
    def all_bullets(c):
        return all(line.lstrip().startswith("- ") for line in c.split("\n"))

    out = []
    for chunk in chunks:
        if out and all_bullets(chunk) and all_bullets(out[-1]):
            out[-1] += "\n" + chunk
        else:
            out.append(chunk)
    return out


def notes_chunk(slide):
    if not slide.has_notes_slide:
        return None
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        return None
    text = (frame.text or "").replace("\x0b", "\n").strip()
    if not text:
        return None
    out = []
    for i, line in enumerate(text.split("\n")):
        line = line.strip()
        prefix = "> 발표자 노트: " if i == 0 else "> "
        out.append((prefix + line).rstrip() if line else ">")
    return "\n".join(out)


def slide_chunks(slide, is_first):
    """한 슬라이드 -> (청크 리스트, 문서 제목 후보 or None, 이미지 수, 표 수, 차트 수)."""
    chunks, images, tables, charts = [], 0, 0, 0
    doc_title = None
    title_shape = slide.shapes.title
    if title_shape is not None and not (
            title_shape.has_text_frame and title_shape.text_frame.text.strip()):
        title_shape = None
    title_done = False

    for kind, shape in walk(slide.shapes):
        if kind == "text":
            lines = para_lines(shape.text_frame)
            if not lines:
                continue
            is_title = (shape is title_shape) or (title_shape is None and not title_done)
            if is_title and not title_done:
                title_done = True
                head, rest = lines[0][1], lines[1:]
                if is_first:
                    doc_title = head           # 1번 슬라이드 첫 줄 = 문서 제목 H1
                    chunks += render_lines(rest)
                else:
                    chunks.append("## " + head)
                    chunks += [t if b else "_%s_" % t for _, t, b in rest]
                continue
            chunks += render_lines(lines)
        elif kind == "table":
            tables += 1
            chunk = table_chunk(shape)
            if chunk:
                chunks.append(chunk)
        elif kind == "chart":
            charts += 1
            images += 1
            chunks.append("[CHART: %s]" % shape.name)
        elif kind == "image":
            images += 1
            chunks.append("[IMAGE: %s]" % shape.name)
        elif kind == "xmltext":
            chunks += ["- " + t for t in a_texts(shape._element)]

    for texts in (a_texts(el) for el in slide.shapes._spTree.iterchildren()
                  if el.tag == MC_ALT):
        if texts:                              # SmartArt 폴백 드로잉
            chunks.append("\n".join("- " + t for t in texts))

    chunks = merge_bullets(chunks)
    note = notes_chunk(slide)
    if note:
        chunks.append(note)
    return chunks, doc_title, images, tables, charts


def convert(path):
    from pptx import Presentation

    prs = Presentation(str(path))
    body, doc_title = [], None
    totals = {"slides": 0, "notes": 0, "images": 0, "tables": 0, "charts": 0,
              "chars": 0, "sparse_slide_numbers": []}
    for i, slide in enumerate(prs.slides, 1):
        chunks, title, images, tables, charts = slide_chunks(slide, i == 1)
        if title and not doc_title:
            doc_title = title
        totals["slides"] = i
        totals["images"] += images
        totals["tables"] += tables
        totals["charts"] += charts
        if notes_chunk(slide):
            totals["notes"] += 1
        chars = len("".join("".join(c.split()) for c in chunks))
        totals["chars"] += chars
        if chars < SPARSE_CHARS:
            totals["sparse_slide_numbers"].append(i)
        body.append("\n\n".join(["<!-- slide %d -->" % i] + chunks))

    totals["sparse_slides"] = len(totals["sparse_slide_numbers"])
    slides = totals["slides"] or 1
    totals["chars_per_slide"] = round(totals["chars"] / slides, 1)
    title = doc_title or path.stem
    return "# %s\n\n%s\n" % (title, "\n\n".join(body)), totals


def main():
    args = [a for a in sys.argv[1:] if a != "--stats"]
    stats_only = "--stats" in sys.argv[1:]
    if len(args) != 1:
        sys.exit("usage: uv run p1-convert.py <file.pptx> [--stats]")
    src = Path(args[0])
    if src.suffix.lower() != ".pptx":
        sys.exit("unsupported extension: %s (expected .pptx)" % src.suffix)
    if not src.is_file():
        sys.exit("no such file: %s" % src)
    try:
        import pptx  # noqa: F401
    except ImportError:
        sys.exit("python-pptx not available — run with: uv run p1-convert.py <file.pptx>")

    md, totals = convert(src)
    if stats_only:
        order = ["slides", "notes", "images", "tables", "charts", "chars",
                 "chars_per_slide", "sparse_slides", "sparse_slide_numbers"]
        print(json.dumps({k: totals[k] for k in order}, ensure_ascii=False))
        return
    if not md.strip() or totals["chars"] == 0:
        sys.exit("empty conversion result — P2(비전 전사) 검토 필요")
    sys.stdout.write(md)


if __name__ == "__main__":
    main()
