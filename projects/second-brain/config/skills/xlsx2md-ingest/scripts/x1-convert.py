#!/usr/bin/env python3
# /// script
# dependencies = ["openpyxl"]
# ///
"""X1: openpyxl XLSX/XLSM -> Clippings MD 변환 랩퍼 (xlsx2md-ingest).

usage: uv run x1-convert.py <file.xlsx|file.xlsm> [--max-rows N] [--sample N] [--formulas]
(uv가 PEP 723 메타데이터로 openpyxl을 자동 준비한다. 순수 Python — Excel 불필요.
 이미 설치된 환경이면 python3로 직접 실행해도 된다.)

md는 stdout으로만 나간다. 시트마다 `## <시트명>` H2로 구분하고, 첫 비어있지 않은
행을 표 헤더로 쓴다. 병합셀은 좌상단 값만 채우고 나머지 칸은 빈칸으로 둔다.
전부 빈 행·열은 건너뛴다.

값 우선 원칙: 수식은 파일에 캐시된 **계산 결과**를 쓴다. 캐시가 없으면(Excel이
한 번도 저장하지 않은 파일) 수식 문자열로 대체하고 stderr에 경고한다.
`--formulas`는 값 뒤에 수식을 `123 (=SUM(A1:A5))`처럼 병기한다.

`--max-rows N`을 넘긴 시트는 X2 모드로 처리한다 — 전체 전사 대신 스키마 +
요약 통계 + 샘플 행만 낸다.

stderr 마지막 줄에 판정·검증용 통계를 출력한다:
  stats sheets=<시트 수> rows=<데이터 행 합계> merged=<병합 범위 수> formulas=<수식 셀 수> truncated=<bool>
"""
import argparse
import datetime as dt
import re
import sys
from pathlib import Path

MAX_CELL = 300  # 셀 텍스트 절단 길이 (md 표 가독성)
CURRENCY = re.compile(r"[₩$€£¥]")


# --- 값 포매팅 ---------------------------------------------------------------

def _num(value, nf: str) -> str:
    """숫자 서식(통화·백분율·천단위)을 보존해 문자열로 만든다."""
    suffix = ""
    # Excel의 두 서식을 구분한다:
    #   0.0%    → 진짜 백분율. 저장값 0.95 를 95.0% 로 표시(×100)
    #   0.0"%"  → 리터럴 % 문자. 저장값 95 를 95.0% 로 표시(배율 없음)
    # 따옴표 밖에 있는 % 만 백분율 서식이다. 이를 구분하지 않으면 후자가 100배로
    # 부풀려진다 (2026-08-26 실측: 저장값 95 → "9500.0%").
    unquoted = re.sub(r'"[^"]*"', "", nf)
    if "%" in unquoted:
        value = value * 100
        suffix = "%"
    elif "%" in nf:
        suffix = "%"
    decimals = None
    if "." in nf:
        frac = nf.split(".", 1)[1]
        decimals = len([c for c in frac if c in "0#"]) or None
    if decimals is None and float(value).is_integer():
        decimals = 0
    grouping = "," if ("#,##" in nf or "0,0" in nf) else ""
    if decimals is None:
        body = format(value, f"{grouping}g") if grouping else format(value, "g")
    else:
        body = format(value, f"{grouping}.{decimals}f")
    symbol = CURRENCY.search(nf)
    return f"{symbol.group(0)}{body}{suffix}" if symbol else f"{body}{suffix}"


def fmt(value, nf: str = "") -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time(0) else value.isoformat(sep=" ")
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, dt.time):
        return value.isoformat()
    if isinstance(value, (int, float)):
        return _num(value, nf or "General")
    return str(value)


def cell_md(text: str) -> str:
    """md 표 셀로 안전하게 만든다 (파이프 이스케이프·개행 접기·절단)."""
    text = text.replace("|", r"\|")
    text = re.sub(r"\s*\n\s*", "<br>", text.strip())
    if len(text) > MAX_CELL:
        text = text[:MAX_CELL].rstrip() + "…"
    return text


# --- 시트 읽기 ---------------------------------------------------------------

def read_matrix(ws, ws_f, use_formulas: bool):
    """시트를 (matrix, raw_matrix, formula_count, no_cache) 로 읽는다.

    matrix: 표시용 문자열 2차원 배열 (빈 행·열 제거 후)
    raw_matrix: 요약 통계용 원시 값 2차원 배열 (같은 모양)
    """
    anchors, covered = {}, set()
    for rng in ws.merged_cells.ranges:
        anchors[(rng.min_row, rng.min_col)] = True
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    covered.add((r, c))

    disp, raw = [], []
    formula_count, no_cache = 0, 0
    for row in ws.iter_rows():
        d_row, r_row = [], []
        for cell in row:
            pos = (cell.row, cell.column)
            if pos in covered:
                d_row.append("")
                r_row.append(None)
                continue
            value = cell.value
            formula = None
            if ws_f is not None:
                fcell = ws_f.cell(row=cell.row, column=cell.column)
                if isinstance(fcell.value, str) and fcell.value.startswith("="):
                    formula = fcell.value
                    formula_count += 1
            text = fmt(value, cell.number_format)
            if formula and not text:
                text = formula  # 캐시된 계산 결과 없음 → 수식 문자열로 대체
                no_cache += 1
            elif formula and use_formulas:
                text = f"{text} ({formula})"
            d_row.append(cell_md(text))
            r_row.append(value)
        disp.append(d_row)
        raw.append(r_row)

    keep_cols = sorted({c for r in disp for c, v in enumerate(r) if v})
    keep_rows = [i for i, r in enumerate(disp) if any(r)]
    matrix = [[disp[i][c] if c < len(disp[i]) else "" for c in keep_cols] for i in keep_rows]
    raw_matrix = [[raw[i][c] if c < len(raw[i]) else None for c in keep_cols] for i in keep_rows]
    return matrix, raw_matrix, formula_count, no_cache


# --- 렌더 --------------------------------------------------------------------

def render_table(matrix) -> str:
    header, body = matrix[0], matrix[1:]
    width = len(header)
    header = [h or f"col{i + 1}" for i, h in enumerate(header)]
    lines = ["| " + " | ".join(header) + " |",
             "|" + "|".join(["---"] * width) + "|"]
    for row in body:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def column_stats(name: str, values) -> str:
    filled = [v for v in values if v is not None and v != ""]
    if not filled:
        return f"| {name} | (빈 열) | 0 | — |"
    nums = [v for v in filled if isinstance(v, (int, float)) and not isinstance(v, bool)]
    dates = [v for v in filled if isinstance(v, (dt.date, dt.datetime))]
    if len(nums) == len(filled):
        kind, detail = "숫자", (f"min {fmt(min(nums))} / max {fmt(max(nums))} / "
                              f"평균 {fmt(round(sum(nums) / len(nums), 2))}")
    elif len(dates) == len(filled):
        kind, detail = "날짜", f"{fmt(min(dates))} ~ {fmt(max(dates))}"
    else:
        uniq = {str(v) for v in filled}
        kind, detail = "텍스트", f"고유값 {len(uniq)}개"
    return f"| {name} | {kind} | {len(filled)} | {cell_md(detail)} |"


def render_summary(matrix, raw_matrix, sample: int) -> str:
    header = [h or f"col{i + 1}" for i, h in enumerate(matrix[0])]
    body, raw_body = matrix[1:], raw_matrix[1:]
    out = ["### 스키마·요약 통계 (X2 — 전량 전사 생략)", "",
           f"- 데이터 행 {len(body):,}행 × {len(header)}열",
           "- 전량 전사 대신 스키마·요약·샘플만 담는다. 상세는 원본 xlsx를 볼 것.",
           "", "| 컬럼 | 타입 | 값 있는 행 | 요약 |", "|---|---|---|---|"]
    for i, name in enumerate(header):
        out.append(column_stats(name, [r[i] for r in raw_body]))
    n = min(sample, len(body))
    if n:
        out += ["", f"### 샘플 행 (상위 {n}행)", "", render_table([matrix[0]] + body[:n])]
    return "\n".join(out)


# --- main --------------------------------------------------------------------

def main() -> None:
    ap = argparse.ArgumentParser(description="XLSX/XLSM -> Clippings MD (stdout)")
    ap.add_argument("src")
    ap.add_argument("--max-rows", type=int, default=0,
                    help="시트 데이터 행이 N을 넘으면 X2(요약)로 처리 (0=제한 없음)")
    ap.add_argument("--sample", type=int, default=5, help="X2 샘플 행 수 (기본 5)")
    ap.add_argument("--formulas", action="store_true", help="값 뒤에 수식 병기")
    args = ap.parse_args()

    src = Path(args.src)
    if src.suffix.lower() not in (".xlsx", ".xlsm"):
        sys.exit(f"unsupported extension: {src.suffix} (expected .xlsx or .xlsm)")
    if not src.is_file():
        sys.exit(f"not found: {src}")
    try:
        import openpyxl
    except ImportError:
        sys.exit("openpyxl not available — run with: uv run x1-convert.py <file.xlsx>")

    wb = openpyxl.load_workbook(src, data_only=True)
    wb_f = openpyxl.load_workbook(src, data_only=False)

    parts, total_rows, total_merged, total_formulas, total_nocache = [], 0, 0, 0, 0
    truncated = False
    for name in wb.sheetnames:
        ws, ws_f = wb[name], wb_f[name]
        merged = len(ws.merged_cells.ranges)
        total_merged += merged
        matrix, raw_matrix, fcount, nocache = read_matrix(ws, ws_f, args.formulas)
        total_formulas += fcount
        total_nocache += nocache
        parts.append(f"## {name}")
        if not matrix:
            parts.append("_(빈 시트 — 값 있는 셀 없음)_")
            continue
        rows = len(matrix) - 1
        total_rows += rows
        notes = [f"{rows:,}행 × {len(matrix[0])}열"]
        if ws.sheet_state != "visible":
            notes.append("숨김 시트")  # 화면에 없어도 내용은 전부 읽힌다 — 게이트 1 확인 대상
        if merged:
            notes.append(f"병합 범위 {merged}개")
        if fcount:
            notes.append(f"수식 셀 {fcount}개")
        parts.append(f"<!-- sheet: {name} — {', '.join(notes)} -->")
        if args.max_rows and rows > args.max_rows:
            truncated = True
            parts.append(render_summary(matrix, raw_matrix, args.sample))
        else:
            parts.append(render_table(matrix))

    print("\n\n".join(parts))
    if total_nocache:
        print(f"warning: 캐시된 계산 결과 없는 수식 셀 {total_nocache}개 — 수식 문자열로 대체함",
              file=sys.stderr)
    print(f"stats sheets={len(wb.sheetnames)} rows={total_rows} merged={total_merged} "
          f"formulas={total_formulas} truncated={str(truncated).lower()}", file=sys.stderr)


if __name__ == "__main__":
    main()
