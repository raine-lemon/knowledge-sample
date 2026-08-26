#!/usr/bin/env python3
# /// script
# dependencies = ["python-docx"]
# ///
"""D1: python-docx DOCX -> MD 변환 랩퍼 (docx2md-ingest).

usage: uv run d1-convert.py <file.docx>           # MD를 stdout으로
       uv run d1-convert.py <file.docx> --stats   # 구조 통계 JSON을 stdout으로

(uv가 PEP 723 메타데이터로 python-docx를 자동 준비한다. python-docx가 이미 설치된
 환경이면 python3로 직접 실행해도 된다.)

문서 순서 보존 — 이 스크립트의 존재 이유:
  `doc.paragraphs`와 `doc.tables`는 각각 **별도의 평면 리스트**라 서로의 위치
  정보를 갖고 있지 않다. 두 리스트를 이어붙이면 표가 전부 문서 맨 끝으로 몰린다
  (2026-08-26 실제 발생 → 수정). 아래 convert()는 `doc.element.body`를 XML 자식
  순서대로 순회해 `w:p`(문단)와 `w:tbl`(표)을 제자리에 섞는다.
  **이 순회 방식을 다른 것으로 바꾸지 말 것.**

추적 변경(track changes)은 '최종본(변경 수락)' 기준으로 읽는다 — 삽입분(w:ins 하위
w:t)은 포함되고, 삭제분은 w:delText 태그라서 자연히 빠진다.
"""
import json
import re
import sys
from pathlib import Path

PSEUDO_HEADING_MAX_CHARS = 80   # D2 판정용: 이보다 짧아야 '서식만으로 만든 제목' 후보
PSEUDO_HEADING_MIN_SZ = 28      # half-point (= 14pt)

Document = qn = Table = Paragraph = None


def load_docx() -> None:
    """python-docx를 로드해 모듈 전역에 바인딩한다 (없으면 안내 후 중단)."""
    global Document, qn, Table, Paragraph
    try:
        from docx import Document as _Document
        from docx.oxml.ns import qn as _qn
        from docx.table import Table as _Table
        from docx.text.paragraph import Paragraph as _Paragraph
    except ImportError:
        sys.exit("python-docx not available — run with: uv run d1-convert.py <file.docx>")
    Document, qn, Table, Paragraph = _Document, _qn, _Table, _Paragraph


# --- 문단 ---------------------------------------------------------------

def para_text(p_el) -> str:
    """w:p 하위를 XML 순서로 훑어 텍스트를 만든다 (각주 참조·이미지 마커 포함)."""
    out = []
    for node in p_el.iter():
        tag = node.tag
        if tag == qn("w:t"):
            out.append(node.text or "")
        elif tag == qn("w:tab"):
            out.append("\t")
        elif tag == qn("w:br"):
            out.append("\n")
        elif tag == qn("w:footnoteReference"):
            out.append("[^fn%s]" % node.get(qn("w:id")))
        elif tag == qn("w:endnoteReference"):
            out.append("[^en%s]" % node.get(qn("w:id")))
        elif tag in (qn("w:drawing"), qn("w:pict")):
            out.append("[IMAGE]")
    return "".join(out).strip()


def list_level(style_lower: str, p_el) -> int:
    """'List Bullet 2' 같은 스타일 접미 숫자와 w:ilvl 중 큰 값을 들여쓰기 깊이로."""
    lvl = 0
    m = re.search(r"(\d)\s*$", style_lower)
    if m:
        lvl = max(int(m.group(1)) - 1, 0)
    ilvl = p_el.find(".//" + qn("w:ilvl"))
    if ilvl is not None:
        try:
            lvl = max(lvl, int(ilvl.get(qn("w:val"))))
        except (TypeError, ValueError):
            pass
    return min(lvl, 5)


def style_prefix(style_name: str, p_el):
    """스타일명 -> (md 접두, 들여쓰기 깊이, 종류). 종류: 'head'|'list'|'text'."""
    s = (style_name or "").strip().lower()
    m = re.match(r"heading\s*(\d)", s)
    if m:
        return "#" * min(int(m.group(1)), 6) + " ", 0, "head"
    if s == "title":
        return "# ", 0, "head"
    if s.startswith("subtitle"):
        return "## ", 0, "head"
    if s.startswith("list bullet"):
        return "- ", list_level(s, p_el), "list"
    if s.startswith("list number"):
        return "1. ", list_level(s, p_el), "list"
    if s.startswith("list paragraph"):
        # 번호 서식 판별에는 numbering 파트 조회가 필요하다 — 불릿으로 통일하고
        # 원문이 번호 목록이면 §4 검증에서 육안 확인한다.
        return "- ", list_level(s, p_el), "list"
    if s.startswith("quote") or s.startswith("intense quote"):
        return "> ", 0, "text"
    return "", 0, "text"


def is_pseudo_heading(p_el, text: str) -> bool:
    """스타일 없이 굵기·크기로만 제목을 흉내낸 문단인지 (D2 라우팅 신호)."""
    if not text or len(text) > PSEUDO_HEADING_MAX_CHARS:
        return False
    total = bold = big = 0
    for r in p_el.findall(".//" + qn("w:r")):
        if not "".join(t.text or "" for t in r.iter(qn("w:t"))).strip():
            continue
        total += 1
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            continue
        b = rpr.find(qn("w:b"))
        if b is not None and b.get(qn("w:val")) not in ("0", "false", "none"):
            bold += 1
        sz = rpr.find(qn("w:sz"))
        if sz is not None:
            try:
                if int(sz.get(qn("w:val"))) >= PSEUDO_HEADING_MIN_SZ:
                    big += 1
            except (TypeError, ValueError):
                pass
    return total > 0 and (bold == total or big == total)


# --- 표 -----------------------------------------------------------------

def table_block(tbl) -> str:
    """w:tbl -> md 표 한 블록. 행 사이에 빈 줄을 넣지 않는다 (넣으면 렌더 깨짐)."""
    rows = []
    try:
        for row in tbl.rows:
            cells = []
            for c in row.cells:
                t = c.text.strip().replace("|", r"\|")
                t = re.sub(r"\s*\n\s*", "<br>", t)
                cells.append(t or " ")
            rows.append(cells)
    except Exception as exc:                      # 비정상 표(병합 이상 등)
        return "<!-- table skipped: %s -->" % exc
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [" "] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |",
             "|" + "|".join(["---"] * width) + "|"]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(lines)


# --- 각주·미주 -----------------------------------------------------------

def collect_notes(doc, kind: str) -> dict:
    """footnotes.xml / endnotes.xml에서 {id: 텍스트}를 모은다."""
    from lxml import etree
    item = qn("w:footnote") if kind == "footnotes" else qn("w:endnote")
    skip = ("separator", "continuationSeparator", "continuationNotice")
    res = {}
    for part in doc.part.package.iter_parts():
        if not str(part.partname).endswith("/%s.xml" % kind):
            continue
        el = getattr(part, "element", None)
        if el is None:
            el = etree.fromstring(part.blob)
        for n in el.findall(item):
            if n.get(qn("w:type")) in skip:
                continue
            txt = " ".join("".join(t.text or "" for t in n.iter(qn("w:t"))).split())
            if txt:
                res[n.get(qn("w:id"))] = txt
    return res


def has_part(doc, suffix: str) -> bool:
    return any(str(p.partname).endswith(suffix) for p in doc.part.package.iter_parts())


# --- 변환 본체 -----------------------------------------------------------

def convert(doc):
    """(md 텍스트, 통계 dict)를 돌려준다."""
    blocks, tables, images, pseudo = [], 0, 0, 0
    heading_styles, paragraphs = {}, 0
    pending_list = []

    def flush_list():
        if pending_list:
            blocks.append("\n".join(pending_list))
            pending_list.clear()

    # ★ 문서 순서(document order) 순회. doc.paragraphs / doc.tables를 쓰면
    #   표가 전부 문서 끝으로 몰린다.
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            paragraphs += 1
            p = Paragraph(child, doc)
            text = para_text(child)
            images += text.count("[IMAGE]")
            if not text:
                continue
            style = (p.style.name if p.style is not None else "") or ""
            prefix, lvl, kind = style_prefix(style, child)
            if kind == "head":
                heading_styles[style] = heading_styles.get(style, 0) + 1
            elif kind == "text" and is_pseudo_heading(child, text):
                pseudo += 1
            if kind == "list":
                pending_list.append("  " * lvl + prefix + text)
            else:
                flush_list()
                blocks.append(prefix + text.replace("\n", "  \n"))
        elif child.tag == qn("w:tbl"):
            flush_list()
            block = table_block(Table(child, doc))
            if block:
                tables += 1
                blocks.append(block)
    flush_list()

    md = "\n\n".join(blocks).strip()

    # 각주·미주 정의를 본문 뒤에 모은다 (본문 참조는 [^fnN]/[^enN]).
    foot, end = collect_notes(doc, "footnotes"), collect_notes(doc, "endnotes")
    used_fn = set(re.findall(r"\[\^fn(\d+)\]", md))
    used_en = set(re.findall(r"\[\^en(\d+)\]", md))
    defs = ["[^fn%s]: %s" % (i, foot[i]) for i in sorted(used_fn, key=int) if i in foot]
    defs += ["[^en%s]: %s" % (i, end[i]) for i in sorted(used_en, key=int) if i in end]
    if defs:
        md += "\n\n" + "\n".join(defs)

    body_xml = doc.element.body
    ins = len(body_xml.findall(".//" + qn("w:ins")))
    dele = len(body_xml.findall(".//" + qn("w:del")))
    comments = len(body_xml.findall(".//" + qn("w:commentRangeStart"))) or (
        1 if has_part(doc, "/comments.xml") else 0)

    stats = {
        "paragraphs": paragraphs,
        "tables": tables,
        "images": images,
        "chars": len("".join(md.split())),
        "heading_styles": heading_styles,
        "has_heading_styles": bool(heading_styles),
        "pseudo_headings": pseudo,
        "footnotes": len(used_fn),
        "endnotes": len(used_en),
        "revisions": {"insertions": ins, "deletions": dele, "comments": comments},
        "has_revisions": bool(ins or dele or comments),
    }
    return md, stats


def main() -> None:
    args = [a for a in sys.argv[1:] if a != "--stats"]
    stats_only = "--stats" in sys.argv[1:]
    if len(args) != 1:
        sys.exit("usage: uv run d1-convert.py <file.docx> [--stats]")
    src = Path(args[0])
    if src.suffix.lower() != ".docx":
        sys.exit("unsupported extension: %s (expected .docx — .doc는 먼저 .docx로 저장)"
                 % src.suffix)
    if not src.is_file():
        sys.exit("no such file: %s" % src)

    load_docx()
    md, stats = convert(Document(str(src)))

    if stats_only:
        print(json.dumps(stats, ensure_ascii=False, indent=2))
        return
    if not md.strip():
        sys.exit("empty conversion result — 텍스트 추출 실패 (D3 검토)")
    sys.stdout.write(md + "\n")


if __name__ == "__main__":
    main()
